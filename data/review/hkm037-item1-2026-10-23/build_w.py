#!/usr/bin/env python3
"""HKM-037 البند 1 — الورقة و: تركيب النص الساري من سلسلة الأدوات.

  python3 build_w.py workbook            # يبني item1-w-decisions.xlsx من compose/*.jsonl
  python3 build_w.py statements          # يحوّل قرارات المالك فيه إلى statements-w.json (إضافة فقط)
"""
import hashlib, json, pathlib, re, sys, warnings
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

warnings.filterwarnings('ignore')
from hijri_converter import Hijri

D = pathlib.Path(__file__).parent
BY = 'owner-inject-hkm037-item1w-2026-10-23'
RUN = 'hkm037-item1w-2026-10-23'
AR = str.maketrans('٠١٢٣٤٥٦٧٨٩', '0123456789')
SHEETS = {'composed': 'و1 نص ساري مركّب', 'already_current': 'و2 ساري أصلًا', 'repealed': 'و3 ملغاة', 'needs_owner': 'و4 يحتاج قرارك'}
ACTION = {'composed': 'نسخة جديدة للمادة بالنص الساري المركّب وتاريخ آخر أداة، ثم تحقق «ساري»',
          'already_current': 'تحقق «ساري» بدليل آخر أداة؛ النص لا يتغير',
          'repealed': 'تحقق «ملغاة» بدليل أداة الإلغاء',
          'needs_owner': 'لا شيء تلقائيًا؛ «موافق مع تصحيح» واكتب النص الكامل في «ملاحظته» لتطبيقه'}
COLS = ['معرّف المادة', 'النظام', 'المادة', 'الحالة', 'الثقة', 'الإجراء عند الموافقة', 'آخر أداة', 'تاريخها', 'النص الساري (مركّب)',
        'مسودة مقترحة', 'نص حكيم الحالي', 'سلسلة التعديلات', 'الملاحظات المغطاة', 'السبب', 'المصدر', 'قرار المالك', 'ملاحظته']
W = [30, 26, 14, 12, 12, 30, 30, 12, 70, 50, 60, 70, 30, 50, 40, 16, 40]


def q(s):
    return 'NULL' if s is None else '$hk$' + str(s).replace('$hk$', '$hk $') + '$hk$'


def hid(*p):
    return hashlib.sha256('|'.join(map(str, p)).encode()).hexdigest()[:24]


def greg(h):
    m = re.search(r'(\d{4})\s*/\s*(\d{1,2})\s*/\s*(\d{1,2})', (h or '').translate(AR))
    if not m:
        return None
    try:
        return Hijri(*map(int, m.groups())).to_gregorian().isoformat()
    except Exception:
        return None


def load():
    res = []
    for f in sorted((D / 'compose').glob('*.jsonl')):
        res += [json.loads(l) for l in open(f, encoding='utf-8') if l.strip()]
    pk = {}
    for f in sorted((D / 'compose').glob('packet*.json')):
        for a in json.load(open(f, encoding='utf-8'))['articles']:
            pk[a['article_id']] = a
    for r in res:
        st, ft = r['status'], r.get('final_text') or ''
        if st == 'incomplete' or r.get('confidence') == 'يحتاج قرار المالك':
            st = 'needs_owner'
        if st == 'composed' and (not ft or '[غير مقروء]' in ft or 'غير محدد في الأداة' in ft or not greg(r.get('last_date_hijri'))):
            st = 'needs_owner'
        r['_st'] = st
        a = pk[r['article_id']]
        opened = [x['note_id'] for x in a['open_notes_readings']]
        cov = [n for n in (r.get('covers_note_ids') or []) if n in opened] or opened
        r['_notes'] = cov
        r['_hk'] = a['hk_text']; r['_url'] = a['law_url']
    return res


def workbook():
    res = load()
    wb = Workbook(); s0 = wb.active; s0.title = 'الملخص'; s0.sheet_view.rightToLeft = True
    YEL = PatternFill('solid', fgColor='FFFF00'); HDR = PatternFill('solid', fgColor='1F3864'); thin = Side(style='thin', color='BFBFBF')
    sh = {}
    for k, name in SHEETS.items():
        ws = wb.create_sheet(name); ws.sheet_view.rightToLeft = True; ws.append(COLS)
        for j, c in enumerate(COLS, 1):
            cell = ws.cell(row=1, column=j); cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF'); cell.fill = HDR
            cell.alignment = Alignment(wrap_text=True, horizontal='center'); ws.column_dimensions[cell.column_letter].width = W[j - 1]
        ws.freeze_panes = 'C2'
        dv = DataValidation(type='list', formula1='"موافق,مرفوض,موافق مع تصحيح,مؤجل"', allow_blank=True); ws.add_data_validation(dv); dv.add('P2:P500')
        sh[k] = ws
    cnt = {k: 0 for k in SHEETS}
    for r in sorted(res, key=lambda x: (x['law'], x['heading'])):
        k = r['_st']; cnt[k] += 1
        chain = '\n'.join(f"{c.get('order')}. {c.get('instrument')} ({c.get('date_hijri') or ''}) — {c.get('effect')}: «{(c.get('quote') or '')[:300]}» [{c.get('source')}]" for c in r.get('chain') or [])
        row = [r['article_id'], r['law'], r['heading'], r['status'], r['confidence'], ACTION[k], r.get('last_instrument'), r.get('last_date_hijri'),
               r.get('final_text') or '', r.get('draft_text') or r.get('partial_text_paras_1_4_not_current') or '', r['_hk'], chain, ' '.join(r['_notes']),
               r.get('reason') or '', r['_url'], '', '']
        ws = sh[k]; ws.append(row); n = ws.max_row
        for j in range(1, len(COLS) + 1):
            c = ws.cell(row=n, column=j); c.font = Font(name='Arial', size=10); c.alignment = Alignment(wrap_text=True, vertical='top')
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        ws.cell(row=n, column=16).fill = YEL; ws.cell(row=n, column=17).fill = YEL
    s0.append(['الورقة و من البند 1 — النص الساري المركّب من سلسلة أدوات كل مادة (HKM-037)'])
    s0.append(['كل نص مبني من نص حكيم الحالي والنصوص الحرفية في صور الأدوات (تشغيل 39987807785) أو نصوص ملاحظات الهيئة، بسلسلة موثّقة لكل مادة. لا يُكتب شيء قبل قرارك.'])
    s0.append([]); s0.append(['الورقة', 'عدد المواد', 'ما يحدث عند «موافق»'])
    for k, name in SHEETS.items():
        s0.append([name, cnt[k], ACTION[k]])
    s0.append(['المجموع', sum(cnt.values()), ''])
    s0.append([]); s0.append(['يكفي أن تقول «موافق على و1 وو2 وو3». الورقة و4 تحتاج نصًّا منك أو اعتماد المسودة المقترحة («موافق مع تصحيح» مع النص في «ملاحظته»).'])
    for row in s0.iter_rows():
        for c in row:
            c.font = Font(name='Arial', size=11, bold=(c.row in (1, 4))); c.alignment = Alignment(wrap_text=True, vertical='top')
    s0.column_dimensions['A'].width = 34; s0.column_dimensions['B'].width = 12; s0.column_dimensions['C'].width = 80
    wb.save(D / 'item1-w-decisions.xlsx')
    print(cnt)


def statements():
    res = {r['article_id']: r for r in load()}
    wb = load_workbook(D / 'item1-w-decisions.xlsx')
    stmts = []; rep = {'articles': 0, 'unit_version': 0, 'verification': 0, 'proposal': 0, 'decision': 0, 'skipped': []}
    for ws in wb.worksheets[1:]:
        hdr = [c.value for c in ws[1]]; ix = {h: i for i, h in enumerate(hdr)}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            dec = row[ix['قرار المالك']]
            if dec not in ('موافق', 'مرفوض', 'موافق مع تصحيح', 'مؤجل'):
                continue
            r = res[row[0]]; aid = r['article_id']; k = r['_st']
            corr = row[ix['ملاحظته']] if dec == 'موافق مع تصحيح' else None
            body = corr or r.get('final_text')
            inst = r.get('last_instrument'); quote = ((r.get('chain') or [{}])[-1].get('quote') or '')[:1500]
            for nid in r['_notes']:
                pid = 'hk37w-' + hid('p', nid)
                stmts.append(f"""INSERT INTO boe.review_proposal (id, note_id, piece_no, category, effect, instrument, instrument_date_hijri, new_text, evidence, confidence, reason, review_run, created_at)
SELECT {q(pid)}, {q(nid)}, 1, 3, {q('تركيب السلسلة: ' + r['status'])}, {q(inst)}, {q(r.get('last_date_hijri'))}, {q(body)}, {q(quote)}, {q(r['confidence'])}, {q(SHEETS[k])}, {q(RUN)}, now()
WHERE NOT EXISTS (SELECT 1 FROM boe.review_proposal WHERE id = {q(pid)})""")
                did = 'hk37wd-' + hid('d', nid, dec)
                stmts.append(f"""INSERT INTO boe.owner_decision (id, proposal_id, decision, corrected_text, owner_note, decided_by, decided_at)
SELECT {q(did)}, {q(pid)}, {q(dec)}, {q(corr)}, {q(row[ix['ملاحظته']])}, 'owner', now()
WHERE NOT EXISTS (SELECT 1 FROM boe.owner_decision WHERE id = {q(did)})""")
                rep['proposal'] += 1; rep['decision'] += 1
            if dec not in ('موافق', 'موافق مع تصحيح'):
                continue
            rep['articles'] += 1
            if k in ('composed', 'needs_owner'):
                vf = greg(r.get('last_date_hijri'))
                if not body or not vf or (k == 'needs_owner' and not corr):
                    rep['skipped'].append(aid); continue
                uid = 'hk37wu-' + hid('u', aid, vf)
                stmts.append(f"""INSERT INTO unit_version (id, unit_id, body, valid_from, evidence_instrument, evidence_url, evidence_quote, created_at)
SELECT {q(uid)}, {q(aid)}, {q(body)}, {q(vf)}::timestamptz, {q(inst)}, {q(r['_url'])}, {q(quote)}, now()
WHERE NOT EXISTS (SELECT 1 FROM unit_version WHERE id = {q(uid)}) AND EXISTS (SELECT 1 FROM legal_articles WHERE id = {q(aid)})""")
                rep['unit_version'] += 1
            status = 'ملغاة' if k == 'repealed' else 'ساري'
            vid = 'vf37w-' + hid('v', aid, status)
            stmts.append(f"""INSERT INTO verification (id, object_type, object_id, verified_status, evidence_instrument, evidence_url, evidence_quote, method, verified_by, verified_at)
SELECT {q(vid)}, 'unit', {q(aid)}, {q(status)}, {q(inst)}, {q(r['_url'])}, {q(quote)}, 'owner_decision', {q(BY)}, now()
WHERE NOT EXISTS (SELECT 1 FROM verification WHERE id = {q(vid)}) AND EXISTS (SELECT 1 FROM legal_articles WHERE id = {q(aid)})""")
            rep['verification'] += 1
    fp = hashlib.md5('\n'.join(stmts).encode()).hexdigest()
    (D / 'statements-w.json').write_text(json.dumps({'by': BY, 'md5': fp, 'count': len(stmts), 'report': rep, 'statements': stmts}, ensure_ascii=False, indent=1), encoding='utf-8')
    print(json.dumps({'md5': fp, 'statements': len(stmts), **{k: (v if not isinstance(v, list) else len(v)) for k, v in rep.items()}}, ensure_ascii=False))


if __name__ == '__main__':
    {'workbook': workbook, 'statements': statements}[sys.argv[1]]()
