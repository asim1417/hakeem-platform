"""ورقة قرار المالك للبند 1 (HKM-037): من نتائج القراء + بيانات الملاحظات."""
import json, pathlib, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

S = pathlib.Path('/tmp/claude-0/-home-user-hakeem-platform/ca620407-2d6b-5011-a62c-22afbdc2c911/scratchpad')
OUT = pathlib.Path('/home/user/hakeem-platform/data/review/hkm037-item1-2026-10-23')
OUT.mkdir(parents=True, exist_ok=True)
notes = {r['note_id']: r for r in json.load(open(S / 'item1_signals.json', encoding='utf-8'))}
res = {}
for i in range(1, 8):
    f = S / 'item1' / f'reader{i}' / 'result.jsonl'
    if f.exists():
        for l in open(f, encoding='utf-8'):
            if l.strip():
                x = json.loads(l)
                res[x['note_id']] = x

TEXT_EFFECTS = {'تعديل مادة', 'تعديل فقرة', 'تعديل جدول', 'تعديل تعريف', 'حذف عبارة', 'إحلال عبارة', 'إضافة فقرة أو مادة', 'دمج أو إعادة ترقيم', 'إلغاء فقرة'}


import re


def compose(n, x):
    """النص الكامل للمادة بعد الأثر، أو None إن تعذّر تركيبه آليًّا بلا لبس."""
    hk = n.get('hk_text') or ''
    e = x.get('effect') or ''
    nt, op, np_ = x.get('new_text'), x.get('old_phrase'), x.get('new_phrase')
    if e == 'تعديل مادة' and nt and '[غير مقروء]' not in nt:
        return nt, 'النص الجديد للمادة كاملًا من الأداة'
    if e == 'إحلال عبارة' and op and np_ and hk:
        c = hk.count(op)
        everywhere = 'أينما' in (x.get('lead_in') or '') or 'أينما' in (n.get('note') or '')
        if c == 1 or (c > 1 and everywhere):
            return hk.replace(op, np_), f'إحلال العبارة ({c} موضع)'
    if e == 'حذف عبارة' and op and hk and hk.count(op) == 1:
        return re.sub(r'[ \t]{2,}', ' ', hk.replace(op, '')), 'حذف العبارة (موضع واحد)'
    return None, None


LATER = re.compile(r'لاحق|أحدث|مرة أخرى|يلزم بعده|يُطبقان فوق|قد يكون لنقص|بصياغة لاحقة|أُلغيت لاحقًا|ألغيت لاحقا')


def sheet_of(n, x):
    if x is None:
        return 'هـ يحتاج قرارك', 'لم تُقرأ'
    if n.get('aid') is None:
        return 'هـ يحتاج قرارك', 'المادة غير موجودة في حكيم'
    if x['status'] == 'not_applicable' and x['confidence'] != 'يحتاج قرار المالك':
        return 'د لا يمس النص', 'تسجيل الأثر دون تغيير النص'
    if x['status'] not in ('found',) or x['confidence'] == 'يحتاج قرار المالك':
        return 'هـ يحتاج قرارك', x.get('reason') or x['status']
    e = x.get('effect') or ''
    if e == 'تفسير أو إجراء لا يمس النص':
        return 'د لا يمس النص', 'تسجيل الأثر دون تغيير النص'
    if e == 'إلغاء مادة':
        return 'ج إلغاء مادة', 'تسجيل المادة «ملغاة»'
    if LATER.search(x.get('reason') or ''):
        return 'و يحتاج تركيبًا', 'القارئ نبّه إلى تعديل لاحق أو سابق غير مطبّق على المادة نفسها؛ تطبيق هذه الأداة وحدها يترك النص غير ساري'
    if x.get('already_in_hakeem'):
        return 'أ ظاهر في حكيم', 'تسجيل التحقق فقط؛ النص لا يتغير'
    body, how = compose(n, x)
    if body:
        return 'ب نص جديد', 'نسخة جديدة للمادة: ' + how
    return 'و يحتاج تركيبًا', 'الأثر مقروء لكن تركيب نص المادة الكامل يحتاج مراجعتك (فقرة، أو عبارة غير موجودة حرفيًا، أو تعديلات سابقة لم تُطبَّق)'


COLS = ['معرّف الملاحظة', 'النظام', 'المادة', 'الفئة', 'الأثر المقترح', 'الإجراء عند الموافقة', 'الأداة', 'تاريخها', 'النص الجديد (حرفي)',
        'العبارة القديمة', 'العبارة الجديدة', 'نص المادة بعد التعديل (مركّب)', 'الدليل (اقتباس من الصورة)', 'ملاحظة الهيئة', 'نص حكيم الحالي', 'المصدر', 'صفحات الصورة',
        'الثقة', 'سبب عدم الحسم', 'قرار المالك', 'ملاحظته']
W = [30, 26, 14, 14, 16, 26, 30, 14, 60, 26, 26, 60, 50, 50, 60, 40, 30, 12, 30, 16, 30]
YEL = PatternFill('solid', fgColor='FFFF00')
HDR = PatternFill('solid', fgColor='1F3864')
thin = Side(style='thin', color='BFBFBF')
F = Font(name='Arial', size=10)
FH = Font(name='Arial', size=10, bold=True, color='FFFFFF')

wb = Workbook()
summ = wb.active
summ.title = 'الملخص'
summ.sheet_view.rightToLeft = True
order = ['أ ظاهر في حكيم', 'ب نص جديد', 'ج إلغاء مادة', 'د لا يمس النص', 'و يحتاج تركيبًا', 'هـ يحتاج قرارك']
sheets = {}
for name in order:
    ws = wb.create_sheet(name)
    ws.sheet_view.rightToLeft = True
    ws.append(COLS)
    for j, c in enumerate(COLS, 1):
        cell = ws.cell(row=1, column=j)
        cell.font = FH; cell.fill = HDR; cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        ws.column_dimensions[cell.column_letter].width = W[j - 1]
    ws.freeze_panes = 'C2'
    dv = DataValidation(type='list', formula1='"موافق,مرفوض,موافق مع تصحيح,مؤجل"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add('T2:T1000')
    sheets[name] = ws

cnt = collections.Counter(); conf = collections.Counter()
rows_out = []
for nid, n in sorted(notes.items(), key=lambda kv: (kv[1]['law'], kv[1]['seq'])):
    x = res.get(nid)
    sh, action = sheet_of(n, x)
    cnt[sh] += 1; conf[(x or {}).get('confidence', 'لم تُقرأ')] += 1
    inst = (x or {}).get('instrument_seen') or ' '.join(filter(None, [n['ikind'], n['ino']]))
    row = [nid, n['law'], n['heading'], sh[0], (x or {}).get('effect') or n['effect'], action, inst, n['idate'] or '',
           (x or {}).get('new_text') or '', (x or {}).get('old_phrase') or '', (x or {}).get('new_phrase') or '',
           (compose(n, x)[0] if (x and sh[0] == 'ب') else ''), (x or {}).get('lead_in') or '',
           n['note'], n['hk_text'] or '', n['url'], ' '.join((x or {}).get('pages') or []), (x or {}).get('confidence') or 'يحتاج قرار المالك',
           (x or {}).get('reason') or ('' if sh[0] != 'هـ' else action), '', '']
    ws = sheets[sh]
    ws.append(row)
    r = ws.max_row
    for j in range(1, len(COLS) + 1):
        c = ws.cell(row=r, column=j)
        c.font = F; c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws.cell(row=r, column=20).fill = YEL; ws.cell(row=r, column=21).fill = YEL
    rows_out.append({'note_id': nid, 'sheet': sh, 'action': action, 'aid': n.get('aid'), 'law_url': n['url'], 'composed': (compose(n, x)[0] if (x and sh[0] == 'ب') else None), 'reading': x})

summ.append(['ورقة قرار البند 1 — ملاحظات التعديل المفتوحة التي لها أداة تعديل مصوّرة (HKM-037)'])
summ.append(['المصدر: صور الأدوات من هيئة الخبراء (تشغيل 39987807785، 736 صورة ببصمات مطابقة)، قراءة بالرؤية حرفيًا. لا شيء يُكتب في القاعدة قبل قرارك.'])
summ.append([])
summ.append(['الورقة', 'العدد', 'ما يحدث عند «موافق»'])
desc = {'أ ظاهر في حكيم': 'سجل تحقق للمادة («ساري») بدليل الأداة؛ النص لا يتغير',
        'ب نص جديد': 'نسخة جديدة للمادة في unit_version بنص الأداة وتاريخها؛ النص القديم يبقى في التاريخ',
        'ج إلغاء مادة': 'سجل تحقق للمادة «ملغاة» بدليل الأداة',
        'د لا يمس النص': 'سجل تحقق «ساري» مع اقتباس التفسير أو الإجراء',
        'و يحتاج تركيبًا': 'لا شيء تلقائيًا؛ يُركَّب نص المادة الكامل يدويًا بعد قرارك ثم يُعرض عليك',
        'هـ يحتاج قرارك': 'لا شيء تلقائيًا؛ يُنفَّذ ما تكتبه في «ملاحظته» أو يُؤجَّل'}
for k in order:
    summ.append([k, cnt[k], desc[k]])
summ.append(['المجموع', sum(cnt.values()), ''])
summ.append([])
summ.append(['الثقة', 'العدد'])
for k, v in conf.most_common():
    summ.append([k, v])
summ.append([])
summ.append(['طريقة القرار: اكتب في عمود «قرار المالك» (أصفر) إحدى القيم: موافق / مرفوض / موافق مع تصحيح / مؤجل. للتصحيح اكتب النص المصحح في «ملاحظته».'])
summ.append(['يكفي أن تقول «موافق على الورقة أ» مثلًا لتطبيق ورقة كاملة.'])
summ.append([])
summ.append(['تنبيهات قبل القرار'])
summ.append(['الورقة ج: 19 مادة من نظام العمل (210–228) تُسجَّل «ملغاة» لأن المرسوم م/1 وتاريخ 1435/1/22هـ (نظام المرافعات الشرعية، أولاً/3) ألغى الباب الرابع عشر «اعتبارًا من إنشاء المحاكم العمالية ومباشرتها اختصاصاتها»؛ انتماء كل مادة إلى الباب مأخوذ من ملاحظة الهيئة لا من الصورة.'])
summ.append(['الورقة و: الأثر مقروء من الصورة، لكن القارئ نبّه إلى أن المادة عُدّلت قبله أو بعده بأداة لم تُطبَّق في حكيم، أو أن التعديل لفقرة يلزم تركيبها؛ تطبيقه وحده يترك نصًّا غير ساري. يُركَّب لكل مادة نصها الساري من سلسلة أدواتها ويُعرض عليك.'])
summ.append(['الورقة هـ: ملاحظات يظهر أنها مرفقة بمادة خاطئة في الهيئة، أو أداتها غير المذكورة، أو المادة غير موجودة في حكيم.'])
for row in summ.iter_rows():
    for c in row:
        c.font = Font(name='Arial', size=11, bold=(c.row in (1, 4))); c.alignment = Alignment(wrap_text=True, vertical='top')
summ.column_dimensions['A'].width = 40; summ.column_dimensions['B'].width = 10; summ.column_dimensions['C'].width = 80
wb.save(OUT / 'item1-decisions.xlsx')
(OUT / 'item1-proposals.json').write_text(json.dumps(rows_out, ensure_ascii=False, indent=1), encoding='utf-8')
print(dict(cnt), dict(conf), 'read', len(res), 'of', len(notes))
