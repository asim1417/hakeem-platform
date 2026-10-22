#!/usr/bin/env python3
"""HKM-030 — يبني مصنف فحص قاعدة الأنظمة من ملفات JSON التي يصدّرها export-audit.ts.
الاستعمال: python3 scripts/hkm-030/build-audit-xlsx.py <json-dir> <out.xlsx>
لا قاعدة ولا شبكة. الخط Arial. الأوراق من اليمين إلى اليسار."""
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = Path(sys.argv[1])
OUT = sys.argv[2]
FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=11)
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, bold=True, size=10)
NOTE = Font(name=FONT, italic=True, size=9, color="555555")


def load(name):
    p = SRC / f"{name}.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else None


def sheet(wb, title, headers, rows, widths=None, note=None, freeze=True):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    r0 = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = NOTE
        r0 = 3
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r0, column=c, value=h)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(rows, r0 + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font = BODY
            cell.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 40)
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = (widths[c - 1] if widths and c - 1 < len(widths) else 18)
    if freeze:
        ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + max(len(rows), 1)}"
    return ws


def pct_sheet_summary(wb, summary, rulings, verif):
    ws = wb.create_sheet("الملخص")
    ws.sheet_view.rightToLeft = True
    ws["A1"] = "فحص قاعدة بيانات الأنظمة — منصة حكيم (HKM-030)"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = f"قراءة من قاعدة الإنتاج (Neon) بتاريخ {str(summary.get('at', ''))[:19]} — كل الأعداد من استعلامات SQL مباشرة؛ النِّسب في العمود D معادلات."
    ws["A2"].font = NOTE
    hdr = ["المؤشر", "القيمة", "الأساس", "النسبة", "ملاحظة"]
    for c, h in enumerate(hdr, 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center")
    s = summary
    rows = [
        ("الأنظمة واللوائح والتنظيمات", s["systems"], None, None, "legal_systems"),
        ("أنظمة بلا مواد (نص كامل في preamble)", s["systems_empty"], s["systems"], "pct", "تُعرض «النص كاملًا»"),
        ("المواد (bis < 100)", s["articles"], None, None, "legal_articles"),
        ("مواد «مكرر»", s["articles_bis"], None, None, ""),
        ("صفوف ليست مواد (bis ≥ 100)", s["non_article_rows"], None, None, "عناوين وجداول"),
        ("مواد مقارنة بنص هيئة الخبراء", s["text_compared"], s["articles"], "pct", "boe.diff"),
        ("منها مطابقة للنص الساري", s["text_same"], s["text_compared"], "pct", "المقياس الرئيس لصحة النص"),
        ("اختلاف متوقع (مادة معدّلة عندنا)", s["text_diff_amended_expected"], s["text_compared"], "pct", ""),
        ("اختلاف وله ملاحظة تعديل غير مطبّقة", s["text_diff_note_unapplied"], s["text_compared"], "pct", "ورقة «ملاحظات الهيئة المفتوحة»"),
        ("اختلاف بلا تفسير", s["text_diff_unexplained"], s["text_compared"], "pct", "ورقة «اختلاف النص»"),
        ("مواد معدّلة (لها نسخة سارية)", s["amended_articles"], s["articles"], "pct", "unit_version"),
        ("نسخ النص", s["versions"], None, None, ""),
        ("ملاحظات الهيئة على موادنا", s["boe_notes_matched"], None, None, ""),
        ("منها معالجة", s["boe_notes_handled"], s["boe_notes_matched"], "pct", ""),
        ("أنظمة بحالة موثقة", s["systems_with_status"], s["systems"], "pct", "verification(work)"),
        ("أنظمة بتاريخ نفاذ", s["systems_with_effective"], s["systems"], "pct", ""),
        ("أنظمة بنوع أداة", s["systems_with_kind"], s["systems"], "pct", ""),
        ("أنظمة بتصنيف", s["systems_with_classification"], s["systems"], "pct", ""),
        ("أنظمة مربوطة بصفحة الهيئة", s["systems_boe_matched"], s["systems"], "pct", "boe.match"),
        ("أنظمة لها صفحة في أم القرى", s["systems_uqn_page"], s["systems"], "pct", "uqn_stage.work"),
        ("أنظمة بلا أي ربط بمصدر", s["systems_no_source"], s["systems"], "pct", ""),
        ("لوائح مربوطة بنظامها الأم", s["regs_linked_to_parent"], None, None, "legal_relations IMPLEMENTS"),
        ("مواد بحالة موثقة", s["articles_with_status"], s["articles"], "pct", ""),
        ("مواد ملغاة", s["articles_repealed"], None, None, ""),
        ("إلغاءات قديمة بلا سند", s["legacy_repeals_unrecorded"], None, None, "ورقة «الإلغاءات القديمة»"),
        ("search_norm ناقص", s["search_norm_missing"], None, None, ""),
        ("مواد بمتجه", s["embedded_articles"], s["articles"], "pct", "embeddings ∪ embedding_version"),
        ("مواد قصيرة جدًا (< 40 حرفًا)", s["very_short_articles"], None, None, "للمراجعة بالعين"),
        ("أنظمة سارية عند الهيئة غير مربوطة", s["boe_active_unmatched"], None, None, "ورقة «الهيئة غير المربوط»"),
        ("هيئة الخبراء: أنظمة / وحدات / مرفقات / صفحات / مرفقات معلّقة", f"{s['boe_laws']} / {s['boe_units']} / {s['boe_attachments']} / {s['boe_attachment_pages']} / {s['boe_attachments_pending']}", None, None, "boe.*"),
        ("أم القرى: أعمال / آثار / وحدات / لقطات / آخر نشر", f"{s['uqn_works']} / {s['uqn_effects']} / {s['uqn_units']} / {s['uqn_snapshots']} / {str(s['uqn_max_published'])[:10]}", None, None, "uqn_stage.*"),
        ("عمليات تعديل أم القرى", s["amendment_ops"], None, None, "uqn_fix.amendment_op"),
        ("الأحكام", s["rulings"], None, None, "judicial_cases"),
        ("أحكام مؤرَّخة", s["rulings_dated"], s["rulings"], "pct", ""),
        ("أجزاء الأحكام / متجهاتها", f"{s['ruling_segments']} / {s['ruling_segment_embeddings']}", None, None, ""),
        ("مبادئ قضائية", s["principles"], None, None, "0 معتمد"),
        ("روابط حكم↔مادة", s["case_article_links"], None, None, "كلها needs_review"),
        ("علاقات مادة↔مادة", s["relations"], None, None, "legal_relations"),
        ("مفاهيم المعجم", s["thesaurus_concepts"], None, None, ""),
    ]
    r = 5
    for label, val, base, kind, note in rows:
        ws.cell(row=r, column=1, value=label).font = BOLD
        ws.cell(row=r, column=2, value=val).font = BODY
        if base is not None:
            ws.cell(row=r, column=3, value=base).font = BODY
            ws.cell(row=r, column=4, value=f"=IF(C{r}=0,0,B{r}/C{r})").number_format = "0.0%"
            ws.cell(row=r, column=4).font = BODY
        ws.cell(row=r, column=5, value=note).font = NOTE
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="التقدير المرجّح لجاهزية القاعدة").font = Font(name=FONT, bold=True, size=12)
    r += 1
    for c, h in enumerate(["البعد", "الوزن", "النسبة", "المساهمة", "الدليل"], 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
    r += 1
    start = r
    dims = [
        ("اكتمال الأنظمة الأساسية", 0.20, f"=1-B{5+28}/{s['boe_laws']}", "أنظمة سارية عند الهيئة غير مربوطة ÷ أنظمة الهيئة"),
        ("صحة النص الساري", 0.25, "=D11", "مطابقة الهيئة"),
        ("التعديلات مطبّقة", 0.15, "=D18", "ملاحظات معالجة"),
        ("الحالة والتواريخ", 0.10, "=(D19+D20)/2", "متوسط الحالة والتاريخ"),
        ("اللوائح وربطها", 0.10, f"=B26/{max(1, s['systems'] - 463)}", "روابط IMPLEMENTS ÷ اللوائح والقواعد (تقريب)"),
        ("الأحكام", 0.10, "=D38*0.8", "مؤرَّخة لكن غير معتمدة"),
        ("الربط حكم↔مادة ومبادئ معتمدة", 0.05, "=0.1", "0 معتمد"),
        ("البحث يعكس النص الساري", 0.05, "=D31", "تغطية المتجهات بعد الفهرسة بالنص الساري"),
    ]
    for label, w, f, ev in dims:
        ws.cell(row=r, column=1, value=label).font = BODY
        ws.cell(row=r, column=2, value=w).number_format = "0%"
        ws.cell(row=r, column=3, value=f).number_format = "0.0%"
        ws.cell(row=r, column=4, value=f"=B{r}*C{r}").number_format = "0.0%"
        ws.cell(row=r, column=5, value=ev).font = NOTE
        r += 1
    ws.cell(row=r, column=1, value="المجموع المرجّح").font = BOLD
    ws.cell(row=r, column=2, value=f"=SUM(B{start}:B{r-1})").number_format = "0%"
    ws.cell(row=r, column=4, value=f"=SUM(D{start}:D{r-1})").number_format = "0.0%"
    ws.cell(row=r, column=4).font = BOLD
    ws.cell(row=r, column=5, value="الأوزان اجتهاد تقرير الفحص 22/10/2026؛ النِّسب معادلات على الجدول أعلاه").font = NOTE
    for c, w in zip("ABCDE", [46, 16, 12, 12, 48]):
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "A5"


def main():
    summary = load("summary")
    systems = load("systems") or []
    diffs = load("text_diffs") or []
    notes = load("open_notes") or []
    repeals = load("legacy_repeals") or []
    boe_un = load("boe_unmatched") or []
    mism = load("count_mismatch") or []
    rulings = load("rulings") or {}
    verif = load("verifications_today") or []

    wb = Workbook()
    wb.remove(wb.active)
    pct_sheet_summary(wb, summary, rulings, verif)

    # الأنظمة
    sheet(wb, "الأنظمة", ["الاسم", "نوع الأداة", "التصنيف", "المجال", "الحالة", "طريقة الحالة", "مسجّل الحالة", "دليل الحالة", "رابط الدليل", "أداة الإصدار", "تاريخ النفاذ", "مواد معلنة", "مواد فعلية", "مواد معدّلة", "مواد بمتجه", "صفحة الهيئة", "حالة الهيئة", "صفحة أم القرى", "نشر أم القرى", "النظام الأم", "مواد مقارنة بالهيئة", "اختلاف بلا تفسير", "ملاحظات مفتوحة", "المعرّف", "الرابط"],
          [[x["name"], x["instrument_kind"], x["classification"], x["domain_title"], x["status"], x["status_method"], x["status_by"], x["status_evidence"], x["status_url"], x["instrument"], x["effective_from"], x["declared_articles"], x["actual_articles"], x["amended_articles"], x["embedded_articles"], x["boe_url"], x["boe_status"], x["uqn_url"], x["uqn_published"], x["parent_law"], x["boe_compared"], x["boe_unexplained"], x["boe_notes_open"], x["id"], f"https://hakeemai.net/law/{x['eli_slug']}" if x.get("eli_slug") else None] for x in systems],
          widths=[48, 12, 28, 22, 12, 10, 24, 40, 30, 36, 12, 8, 8, 8, 8, 30, 12, 30, 12, 36, 10, 10, 10, 26, 40],
          note="كل الأنظمة (1 صف لكل نظام). الحالة = آخر سجل في verification. استعمل الفلاتر: الحالة فارغة = بلا حالة؛ تاريخ النفاذ فارغ = بلا تاريخ؛ التصنيف فارغ = بلا تصنيف.")

    # النواقص
    gaps = [
        ("1", "تعديلات هيئة الخبراء غير مطبّقة على موادنا", summary["boe_notes_matched"] - summary["boe_notes_handled"], "زحف R1 (مرفقات الهيئة) ثم حقن بأداة hkm30_compose؛ 279 منها بنصها جاهزة بعد الاعتماد", "ورقة «ملاحظات الهيئة المفتوحة»"),
        ("2", "أنظمة سارية عند الهيئة غير موجودة عندنا", 6, "زحف R0 (6 أنظمة)؛ الـ4 الأخرى طبعات ثانية لأنظمة مربوطة", "ورقة «الهيئة غير المربوط»"),
        ("3", "اختلاف نص عن الهيئة بلا تفسير", summary["text_diff_unexplained"], "≥0.95 متطابق الحكم (قرار المالك: نصنا المعتمد)؛ 0.3–0.95 مراجعة بالعين؛ <0.3 مطابقة بالمحتوى", "ورقة «اختلاف النص»"),
        ("4", "تعارض عدد المواد مع الهيئة", len(mism), "مطابقة بالمحتوى لا بالرقم", "ورقة «تعارض العدد»"),
        ("5", "إلغاءات قديمة بلا سند", summary["legacy_repeals_unrecorded"], "زحف R5 (م/191 وأداتان بلا مرجع)", "ورقة «الإلغاءات القديمة»"),
        ("6", "أنظمة بلا حالة", summary["systems"] - summary["systems_with_status"], "13 لوزارة العدل (R6) + 13 طبعات قديمة (R2-b)", "ورقة «الأنظمة» (الحالة فارغة)"),
        ("7", "أنظمة بلا تاريخ نفاذ", summary["systems"] - summary["systems_with_effective"], "8 خليجية «وفق الإجراءات الدستورية»؛ الباقي بلا صفحة", "ورقة «الأنظمة»"),
        ("8", "أنظمة بلا تصنيف", summary["systems"] - summary["systems_with_classification"], "تصنيف يدوي أو من صفحة الهيئة؛ تصنيف أم القرى عام", "ورقة «الأنظمة»"),
        ("9", "أنظمة بلا أي ربط بمصدر", summary["systems_no_source"], "زحف R2-b", "ورقة «الأنظمة» (صفحتا الهيئة وأم القرى فارغتان)"),
        ("10", "أدوات بعينها", 6, "قرار 419، تسمية وزارة النقل، الآثار، مرسوم حقوق المؤلف، م/89، مرسوم قرار 154 (R3/R4)", "أمر الزحف"),
        ("11", "مواد قصيرة جدًا", summary["very_short_articles"], "مراجعة بالعين (أغلبها «تُلغى» أو عناوين)", ""),
        ("12", "الأحكام غير معتمدة", summary["rulings"], "خط اعتماد؛ تنوع ضعيف (تجارية 92%)", "ورقة «الأحكام»"),
        ("13", "روابط حكم↔مادة ومبادئ بلا اعتماد", summary["case_article_links"] + summary["principles"], "خط مراجعة آلي + عينة بشرية", ""),
        ("14", "OpenSearch يرفض الكتابة (403)", 1, "فحص حساب العنقود أو أسراره", ""),
        ("15", "boe.diff للأنظمة الأربعة المربوطة اليوم لم يُحسب", 4, "تشغيل مقارِن الهيئة", ""),
    ]
    sheet(wb, "النواقص", ["#", "الخلل", "الحجم", "ما يلزمه", "التفصيل"], [list(g) for g in gaps], widths=[5, 48, 10, 70, 36],
          note="ما بقي من خلل بعد دفعات 18–22 أكتوبر 2026، مرتبًا بالأثر.")

    def cls(sim):
        try:
            v = float(sim)
        except (TypeError, ValueError):
            return ""
        return "≥0.95 متطابق الحكم" if v >= 0.95 else ("0.9–0.95" if v >= 0.9 else ("0.3–0.9 مراجعة" if v >= 0.3 else "<0.3 نص مختلف"))

    sheet(wb, "اختلاف النص", ["النظام", "المادة", "مكرر", "حكم الهيئة", "حالة الهيئة", "التشابه", "الفئة", "طول نصنا", "طول نص الهيئة", "نصنا (مقتطف)", "نص الهيئة (مقتطف)", "صفحة الهيئة", "معرّف المادة"],
          [[d["law"], d["article"], d["bis"], d["boe_verdict"], d["boe_status"], float(d["similarity"]) if d["similarity"] is not None else None, cls(d["similarity"]), d["our_len"], d["boe_len"], d["our_text"], d["boe_text"], d["boe_url"], d["article_id"]] for d in diffs],
          widths=[36, 7, 6, 18, 10, 9, 18, 8, 8, 60, 60, 30, 26],
          note="مواد نصها الساري يختلف عن نص هيئة الخبراء بلا نسخة ولا ملاحظة تعديل. التشابه بعد التطبيع (hkm30_norm).")

    sheet(wb, "ملاحظات الهيئة المفتوحة", ["النظام", "المادة", "مكرر", "الأثر", "الأداة", "بنص جديد", "تحتاج مراجعة", "الملاحظة (مقتطف)", "صفحة الهيئة", "معرّف المادة"],
          [[n["law"], n["article"], n["bis"], n["effect_type"], n["instrument"], "نعم" if n["has_new_text"] else "", "نعم" if n.get("needs_review") else "", n["note"], n["boe_url"], n["article_id"]] for n in notes],
          widths=[36, 7, 6, 14, 30, 8, 8, 80, 30, 26],
          note="ملاحظات تعديل من هيئة الخبراء على مواد عندنا لم تُطبَّق بعد (لا نسخة ولا تحقق). «بنص جديد» = يمكن حقنها آليًّا بعد الاعتماد.")

    sheet(wb, "الإلغاءات القديمة", ["النظام", "المادة", "الأداة", "التاريخ الهجري", "المصدر", "معرّف المادة"],
          [[r["law"], r["article"], r["instrument"], r["hijri_date"], r["source"], r["article_id"]] for r in repeals],
          widths=[36, 7, 40, 14, 16, 26],
          note="إلغاءات في article_amendments القديم بلا نسخة ولا تحقق «ملغاة». 203 منها مغطاة بحالة «مستبدل» للنظام كله؛ الباقي يحتاج سند الأداة (R5).")

    sheet(wb, "الهيئة غير المربوط", ["النظام عند الهيئة", "الحالة", "المواد", "الإصدار", "النشر", "أقرب نظام عندنا", "التشابه", "الرابط", "guid"],
          [[b["name"], b["status"], b["article_count"], b["issued_hijri"], b["published_hijri"], b["nearest_hakeem"], float(b["nearest_sim"]) if b["nearest_sim"] is not None else None, b["url"], b["guid"]] for b in boe_un],
          widths=[48, 10, 8, 14, 14, 48, 8, 40, 36],
          note="أنظمة في بوابة هيئة الخبراء غير مربوطة بأي نظام عندنا. التشابه ≥ 0.8 غالبًا طبعة ثانية لنظام مربوط؛ الساري منها بلا نظير = ناقص فعلًا (R0).")

    sheet(wb, "تعارض العدد", ["النظام عندنا", "موادنا", "اسم الهيئة", "مواد الهيئة", "حالة الهيئة", "الرابط"],
          [[m["name"], m["ours"], m["boe_name"], m["boe"], m["boe_status"], m["url"]] for m in mism],
          widths=[48, 8, 48, 8, 10, 40], note="أنظمة مربوطة وعدد موادها يختلف عن عدد الهيئة؛ المطابقة بالرقم غير صالحة لها.")

    rr = []
    for k in ("by_court", "by_degree", "review", "principles", "links", "segments"):
        for t in rulings.get(k) or []:
            key = next(iter(t.values()))
            rr.append([k, str(key), t["n"]])
    for t in rulings.get("by_year") or []:
        rr.append(["by_year", str(t["year_h"]), t["n"]])
    rr.insert(0, ["total", "الأحكام", rulings.get("total")])
    rr.insert(1, ["dated", "مؤرَّخة", rulings.get("dated")])
    rr.insert(2, ["with_appeal", "فيها نص استئناف", rulings.get("with_appeal")])
    sheet(wb, "الأحكام", ["البعد", "القيمة", "العدد"], rr, widths=[16, 30, 10], note="توزيعات الأحكام (judicial_cases / ruling_meta / ruling_segments / judicial_principles / legal_article_case_links).")

    sheet(wb, "سجلات التحقق", ["مسجّل التحقق", "النوع", "الحالة", "الطريقة", "العدد", "منذ"],
          [[v["verified_by"], v["object_type"], v["verified_status"], v["method"], v["n"], v["since"]] for v in verif],
          widths=[40, 8, 24, 14, 8, 12], note="كل دفعات التحقق في جدول verification (إضافة فقط).")

    batches = [
        ("2026-10-18", "إحلال «البنك المركزي السعودي» (م/36)", "129 نسخة", "76077aaf", "br-weathered-darkness-aqpxfgln"),
        ("2026-10-18", "search_norm من النص الساري", "582", "1b8ef25c", "br-square-bar-aqtdhyq2"),
        ("2026-10-18", "الآثار «الجهة المختصة» (قرار 1012)", "4 نسخ", "76040e3f", "br-dry-bird-aqwb3t7p"),
        ("2026-10-18", "الأساس المتقادم (تبادل المنافع، المرور، الصادرات)", "8 نسخ", "20e6b8e1", "br-late-night-aqwk29lt"),
        ("2026-10-19", "حذف العنوان من داخل النص", "744", "37c30192", "br-delicate-base-aq9lpq93"),
        ("2026-10-19", "المرحلتان 6 و7 (11 مادة، 19 تحقق، 59 تاريخ، غسل الأموال)", "90", "50b9656c", "br-floral-poetry-aqltfke2"),
        ("2026-10-22", "نوع الأداة، التصنيف، روابط IMPLEMENTS، articleCount", "1,155 / 74 / 62 / 13", "f03368a5", "br-calm-thunder-aqn4raj0"),
        ("2026-10-22", "تاريخ الحكم من ruling_meta", "51,044", "6788b4cf", "br-calm-thunder-aqn4raj0"),
        ("2026-10-22", "تاريخ النفاذ (النشر/الموافقة)", "433", "91fff3dd", "br-misty-lab-aq4azblk"),
        ("2026-10-22", "أداة الاستبدال للأنظمة المستبدلة", "179 تحقق", "7fa39568", "br-super-meadow-aqxqd950"),
        ("2026-10-22", "المتجهات لكل ناقص وقديم بالنص الساري", "12,670", "التشغيل 39787828349", "embedding_version (إضافة فقط)"),
        ("2026-10-22", "ربط 4 أنظمة بصفحة الهيئة (boe.match)", "4", "53dff99d", "br-super-meadow-aqxqd950"),
        ("2026-10-22", "قاعدة «النشر دليل السريان»", "606 + 1 + 8 + 2 تحقق", "6380f43b / af3cc2d7 / de90dd6a / 03f902e6", "br-wild-wind-aq20xnh3"),
    ]
    sheet(wb, "الدفعات المطبّقة", ["التاريخ", "الدفعة", "الصفوف", "البصمة", "نقطة الاسترجاع"], [list(b) for b in batches], widths=[12, 60, 24, 36, 32],
          note="كل دفعة: بروفة على فرع تجريبي ← نقطة استرجاع ← المعاملة نفسها على الإنتاج ← البصمة تطابق البروفة ← سطر في docs/PROJECT_STATUS.md.")

    gates = [
        ("تغطية", "0 نظام ساري عند الهيئة ناقص؛ 0 نظام بلا ربط بمصدر", "6 ناقصة؛ 46 بلا ربط"),
        ("نص", "مطابقة ≥ 99% و0 اختلاف بلا تفسير", f"{summary['text_same']}/{summary['text_compared']}؛ {summary['text_diff_unexplained']} بلا تفسير"),
        ("تعديلات", "0 ملاحظة تعديل مفتوحة", f"{summary['boe_notes_matched'] - summary['boe_notes_handled']} مفتوحة"),
        ("حالة وتاريخ", "100% حالة و100% تاريخ نفاذ", f"{summary['systems_with_status']} / {summary['systems_with_effective']} من {summary['systems']}"),
        ("لوائح", "كل لائحة مربوطة ومصنفة ونوعها محدد", f"{summary['regs_linked_to_parent']} مربوطة؛ {summary['systems'] - summary['systems_with_classification']} بلا تصنيف"),
        ("بحث", "مصدر نصي واحد؛ 0 متجه قديم؛ «مكرر» والحالة في كل نتيجة", f"المتجهات {summary['embedded_articles']}/{summary['articles']}؛ OpenSearch معطّل (403)"),
        ("أحكام", "تاريخ ودرجة لكل حكم؛ خط اعتماد", f"{summary['rulings_dated']}/{summary['rulings']} مؤرَّخ؛ 0 معتمد"),
        ("ربط", "روابط ومبادئ معتمدة ≥ 90% من المراجَع", "0 معتمد"),
    ]
    sheet(wb, "بوابات الجاهزية", ["البوابة", "شرط القبول", "الوضع الآن"], [list(g) for g in gates], widths=[14, 60, 50],
          note="متى نقول إن القاعدة جاهزة من الألف إلى الياء: حين تتحقق البوابات الثماني كلها.")

    crawl = [
        ("R0", "6 أنظمة سارية ناقصة (هيئة الخبراء)", "6"),
        ("R1", "مرفقات الهيئة المعلّقة (أدوات الإصدار أولًا)", str(summary["boe_attachments_pending"])),
        ("R2-b", "أنظمة بلا صفحة معروفة (طبعات قديمة)", "13"),
        ("R6", "لوائح وقواعد وزارة العدل (laws.moj.gov.sa)", "13"),
        ("R5", "سندات الإلغاءات: م/191 وأداتان بلا مرجع", "3"),
        ("R3", "ما نُشر في أم القرى بعد 2026-09-18 (ومنه مرسوم قرار 154)", "—"),
        ("R4", "قرار 419، تسمية وزارة النقل، الآثار، مرسوم حقوق المؤلف، م/89", "5"),
    ]
    sheet(wb, "أمر الزحف", ["الطلب", "الهدف", "العدد"], [list(c) for c in crawl], widths=[8, 70, 10],
          note="ملخص أمر الزحف 2.2 (الملف الكامل سُلِّم للمالك). المصادر: laws.boe.gov.sa، www.uqn.gov.sa، laws.moj.gov.sa.")

    wb.save(OUT)
    print(f"saved {OUT}: sheets={len(wb.sheetnames)} systems={len(systems)} diffs={len(diffs)} notes={len(notes)} repeals={len(repeals)}")


if __name__ == "__main__":
    main()
