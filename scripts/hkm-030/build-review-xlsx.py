#!/usr/bin/env python3
"""HKM-030 دفعة 22/10 (ب) — مصنف المراجعة البشرية لما نُفِّذ آليًّا وما بقي.
python3 scripts/hkm-030/build-review-xlsx.py <scratch-dir> <out.xlsx>
المدخلات: batch2-trial-results.json (الذيول المحذوفة)، textdiff_live.json (الاختلافات الباقية)، statements.json (التصنيفات والمطابقات).
"""
import json, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC, OUT = Path(sys.argv[1]), sys.argv[2]
REPO = Path(__file__).resolve().parents[2]
FONT = "Arial"
HDR = Font(name=FONT, bold=True, color="FFFFFF", size=10)
FILL = PatternFill("solid", fgColor="1F4E79")
BODY = Font(name=FONT, size=10)
NOTE = Font(name=FONT, italic=True, size=9, color="555555")


def sheet(wb, title, headers, rows, widths, note):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.cell(row=1, column=1, value=note).font = NOTE
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font, cell.fill = HDR, FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, row in enumerate(rows, 4):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font = BODY
            cell.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 40)
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{3 + max(len(rows), 1)}"
    return ws


wb = Workbook()
wb.remove(wb.active)
trial = json.load(open(SRC / "batch2-trial-results.json", encoding="utf-8"))
diffs = json.load(open(SRC / "textdiff_live.json", encoding="utf-8"))
st = json.load(open(REPO / "data/renumber/hkm030-batch2-2026-10-22/statements.json", encoding="utf-8"))

# 1) الذيول المحذوفة (نُفِّذ — للمراجعة بعد التنفيذ)
suf = trial["groups"]["g3_tail"]["report"]["suffixes"]
kind = lambda s: "تذييل موقع/حقوق" if ("حقوق محفوظة" in s or "عرض كل المواد" in s) else ("نص مادة تالية" if "المادة" in s else "عنوان باب/فصل")
sheet(wb, "الذيول المحذوفة (نُفّذ)", ["النظام", "المادة", "النوع", "الذيل المحذوف", "قرار المراجع"],
      [(r["law"], r["an"], kind(r["suffix"]), r["suffix"].strip(), "") for r in suf], [44, 8, 16, 70, 18],
      "حُذف من آخر المتن ما بعد نهاية نص الهيئة (بعد التطبيع). النص السابق محفوظ في legal_text_provenance (source_id = hkm030-tail-strip-2026-10-22). اكتب «رجوع» لو أردت إعادة أي مادة.")

# 2) الاختلافات الباقية للمراجعة (لم يُنفَّذ)
rem = [r for r in diffs if r["k"] != "trailing"]
sheet(wb, "اختلاف صياغة (للمراجعة)", ["النظام", "المادة", "النوع", "التشابه", "طول نصنا", "طول الهيئة", "بداية نصنا", "بداية نص الهيئة", "آخر نصنا", "آخر نص الهيئة", "القرار (نصنا/الهيئة)"],
      [(r["law"], r["an"], {"w08": "صياغة ≥0.8", "w05": "صياغة 0.5–0.8", "ours_contains_boe": "نصنا يحوي نص الهيئة وزيادة في أوله", "boe_contains_ours": "نص الهيئة أطول (نصنا مقطوع؟)"}.get(r["k"], r["k"]), r["sim"], r["ol"], r["bl"], r["our_head"], r["boe_head"], r["our_tail"], r["boe_tail"], "") for r in rem],
      [40, 8, 22, 9, 9, 9, 40, 40, 40, 40, 16],
      "اختلافات لا تُحسم آليًّا: أغلبها ترقيم بنود (1- 2-) حذفه عارض الهيئة، أو عنوان موضوعي في أول المادة عندنا، أو تشكيل. تقرير بشري: أي مصدر يتقدم؟")

# 3) التصنيف بالقاعدة (نُفِّذ)
sheet(wb, "التصنيف بالقاعدة (نُفّذ)", ["المعرّف", "النظام", "التصنيف المعطى", "هامش الثقة", "قرار المراجع"],
      [(r["id"], r["name"], r["classification"], r["margin"], "") for r in st["rule_classifications"]], [28, 60, 40, 10, 18],
      "تصنيف من ألفاظ الاسم بالمقارنة مع 587 نظامًا مصنفًا (دقة 94.7% في اختبار الاستبعاد الواحد عند هامش ≥ 1.5). و25 نظامًا آخر أخذ تصنيف النظام الأم الوارد اسمه في اسمه (لا يحتاج مراجعة).")

# 4) مطابقات الهيئة المصححة (نُفِّذ)
rows = []
for s in st["groups"]["g1_match"]:
    if ":" in s["name"] and s["name"].split(":")[0] in ("repoint", "unlink", "link"):
        k, d = s["name"].split(":", 1)
        rows.append(({"repoint": "تحويل الصفحة إلى سجل الإصدار الصحيح", "unlink": "فكّ ربط صفحة إصدار ملغى", "link": "ربط جديد بالمحتوى"}[k], d))
for L in st["groups"]["g2_ingest"]:
    if L["name"].startswith("system:"):
        rows.append(("نظام جديد من لقطة الهيئة", L["name"].split(":", 1)[1]))
    if L["name"].startswith("supersede:"):
        rows.append(("تعليم القديم «مستبدل»", L["name"].split(":", 1)[1]))
sheet(wb, "المطابقات والأنظمة الجديدة", ["الإجراء", "النظام", "ملاحظة المراجع"], [(a, b, "") for a, b in rows], [34, 70, 24],
      "كل الصفوف في boe.match إضافة (المخطط append-only؛ boe.current_match يأخذ الأحدث). الأنظمة الثمانية نصّها منسوخ حرفيًّا من boe.unit وبصمته مطابقة.")

wb.save(OUT)
print("saved", OUT, {ws.title: ws.max_row - 3 for ws in wb.worksheets})
