"""
export_systems_excel.py — تصدير قاعدة الأنظمة وتواريخ إصدارها إلى ملف Excel.

المصادر (بلا اختلاق — التاريخ يُترك فارغًا إن لم يرد في مصدر موثّق):
  • data/legal_systems_classified.json  ← قائمة الأنظمة (485) + الرمز + المجال + عدد المواد
  • lib/modules/legal-core/core-systems.ts ← تواريخ الإصدار المتحقَّق منها من بوابة وزارة العدل
  • data/moj-regs.json + data/moj-instruments.json ← اللوائح والأدوات الإجرائية (تاريخ إصدار رسمي)
  • اسم النظام نفسه إن تضمّن سنة الإصدار (مثل «نظام البريد 1406هـ») ← سنة فقط

المخرج: reports/saudi-systems-issuance.xlsx
التشغيل: python3 scripts/export_systems_excel.py   (يتطلب openpyxl و hijridate)
"""
import json
import re
from datetime import date
from pathlib import Path

from hijridate import Hijri
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
OUT = ROOT / "reports" / "saudi-systems-issuance.xlsx"

PROCEDURAL_DOMAINS = {"procedure", "enforcement", "evidence", "arbitration", "notarization"}

SRC_CORE = "بوابة وزارة العدل (تحقق رسمي)"
SRC_MOJ = "بوابة وزارة العدل (لائحة/أداة رسمية)"
SRC_NAME = "سنة واردة في اسم النظام"


def norm(s: str) -> str:
    s = re.sub(r"[ً-ْٰـ]", "", s or "")
    s = re.sub(r"[إأآ]", "ا", s).replace("ى", "ي").replace("ة", "ه")
    return re.sub(r"\s+", " ", s).strip()


def to_gregorian(h: str | None) -> date | None:
    if not h:
        return None
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})$", h)
    if not m:
        return None
    try:
        return Hijri(int(m[1]), int(m[2]), int(m[3])).to_gregorian()
    except (ValueError, OverflowError):
        return None


def fmt_hijri(h: str | None) -> str:
    if not h:
        return ""
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})$", h)
    return f"{int(m[3])}/{int(m[2])}/{m[1]}هـ" if m else f"{h}هـ"


def instrument_type(name: str) -> str:
    first = name.split()[0].replace("ال", "", 1) if name.startswith("ال") else name.split()[0]
    return {
        "نظام": "نظام", "تنظيم": "تنظيم", "لائحة": "لائحة", "لوائح": "لائحة",
        "ترتيبات": "ترتيبات", "ترتيب": "ترتيبات", "قانون": "قانون",
        "قواعد": "قواعد", "ضوابط": "ضوابط", "أدلة": "دليل", "آلية": "آلية",
    }.get(first, "أخرى")


def load_core() -> dict[str, dict]:
    src = (ROOT / "lib/modules/legal-core/core-systems.ts").read_text(encoding="utf-8")
    out = {}
    for m in re.finditer(r'\{\s*lawName:\s*"([^"]+)".*?issuanceDateH:\s*(null|"([^"]+)").*?serial:\s*(null|"([^"]+)")', src):
        out[norm(m[1])] = {"date": m[3], "serial": m[5]}
    return out


def main() -> None:
    systems = json.loads((DATA / "legal_systems_classified.json").read_text(encoding="utf-8"))["systems"]
    core = load_core()

    rows = []
    for s in sorted(systems, key=lambda x: x["order"]):
        name = s["name"]
        c = core.get(norm(name))
        date_h, source, url, year_only = None, "", "", None
        if c and c["date"]:
            date_h, source = c["date"], SRC_CORE
            url = f"https://laws.moj.gov.sa/ar/legislation/{c['serial']}" if c["serial"] else ""
        else:
            ym = re.search(r"(1[34]\d\d)\s*هـ", name)
            if ym:
                year_only, source = ym[1], SRC_NAME
        rows.append({
            "code": s["code"], "name": name, "type": instrument_type(name),
            "domain": s["domainTitle"], "nature": "إجرائي" if s["domain"] in PROCEDURAL_DOMAINS else "موضوعي/تنظيمي",
            "parent": "", "articles": s["articleCount"],
            "date_h": date_h, "year_h": year_only, "source": source, "url": url,
        })

    seq = 0
    for f in ("moj-regs.json", "moj-instruments.json"):
        for item in json.loads((DATA / f).read_text(encoding="utf-8")):
            sy = item["system"]
            seq += 1
            rows.append({
                "code": f"MOJ-{seq:03d}", "name": sy["name"], "type": instrument_type(sy["name"]),
                "domain": sy.get("classification") or "", "nature": "إجرائي",
                "parent": sy.get("parentSystem") or "", "articles": sy.get("officialArticleCount"),
                "date_h": sy.get("issuanceDateH"), "year_h": None, "source": SRC_MOJ,
                "url": sy.get("sourceUrl") or "",
            })

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=11)
    body_font = Font(name="Arial", size=10)
    thin = Side(style="thin", color="C9CED6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    dated_fill = PatternFill("solid", fgColor="E8F3EA")
    year_fill = PatternFill("solid", fgColor="FFF6DA")

    def sheet(ws, headers, data, widths):
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for cell in ws[1]:
            cell.fill, cell.font, cell.border = header_fill, header_font, border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r in data:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font, cell.border = body_font, border
                cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=False)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 30

    # ── الورقة 1: كل الأنظمة ──
    ws = wb.active
    ws.title = "الأنظمة"
    headers = ["م", "الرمز", "اسم النظام / الأداة", "النوع", "المجال", "الطبيعة", "النظام الأم",
               "عدد المواد", "تاريخ الإصدار (هـ)", "ما يوافقه (م)", "سنة الإصدار (هـ)", "مصدر التاريخ", "رابط المصدر الرسمي"]
    data = []
    for i, r in enumerate(rows, 1):
        g = to_gregorian(r["date_h"])
        year = r["date_h"][:4] if r["date_h"] else r["year_h"]
        data.append([i, r["code"], r["name"], r["type"], r["domain"], r["nature"], r["parent"],
                     r["articles"], fmt_hijri(r["date_h"]), g, int(year) if year else None,
                     r["source"] or "غير متوفر في القاعدة", r["url"]])
    sheet(ws, headers, data, [6, 11, 58, 10, 24, 14, 26, 10, 16, 14, 12, 32, 52])
    for row in ws.iter_rows(min_row=2):
        row[9].number_format = "yyyy-mm-dd"
        src = row[11].value
        if src in (SRC_CORE, SRC_MOJ):
            for cell in row:
                cell.fill = dated_fill
        elif src == SRC_NAME:
            for cell in row:
                cell.fill = year_fill
        if row[12].value:
            row[12].hyperlink = row[12].value
            row[12].font = Font(name="Arial", size=10, color="1F5FBF", underline="single")

    # ── الورقة 2: المؤرَّخ فقط (مرتب زمنياً) ──
    ws2 = wb.create_sheet("المؤرخة")
    dated = [r for r in rows if r["date_h"] or r["year_h"]]
    dated.sort(key=lambda r: (r["date_h"] or f"{r['year_h']}-99-99"))
    data2 = []
    for i, r in enumerate(dated, 1):
        data2.append([i, r["name"], r["type"], r["nature"], r["parent"], fmt_hijri(r["date_h"]) or f"{r['year_h']}هـ",
                      to_gregorian(r["date_h"]), r["source"]])
    sheet(ws2, ["م", "اسم النظام / الأداة", "النوع", "الطبيعة", "النظام الأم", "تاريخ الإصدار (هـ)", "ما يوافقه (م)", "المصدر"],
          data2, [6, 58, 10, 14, 26, 16, 14, 32])
    for row in ws2.iter_rows(min_row=2):
        row[6].number_format = "yyyy-mm-dd"

    # ── الورقة 3: ملخص ──
    ws3 = wb.create_sheet("ملخص")
    n_core = sum(r["source"] == SRC_CORE for r in rows)
    n_moj = sum(r["source"] == SRC_MOJ for r in rows)
    n_name = sum(r["source"] == SRC_NAME for r in rows)
    summary = [
        ["إجمالي السجلات", len(rows)],
        ["أنظمة القاعدة", len(systems)],
        ["لوائح وأدوات إجرائية (وزارة العدل)", seq],
        ["الأنظمة الإجرائية", sum(r["nature"] == "إجرائي" for r in rows)],
        ["بتاريخ إصدار كامل موثّق — أنظمة أساسية", n_core],
        ["بتاريخ إصدار كامل موثّق — لوائح وأدوات", n_moj],
        ["بسنة إصدار من اسم النظام فقط", n_name],
        ["بلا تاريخ في القاعدة", len(rows) - n_core - n_moj - n_name],
        ["تاريخ التصدير", date.today().isoformat()],
        ["ملاحظة", "لم يُضَف أي تاريخ غير وارد في مصدر موثّق. تحويل الهجري للميلادي وفق تقويم أم القرى."],
    ]
    sheet(ws3, ["البند", "القيمة"], summary, [42, 80])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"✓ {OUT.relative_to(ROOT)} — {len(rows)} سجل (مؤرخ كامل: {n_core + n_moj}، سنة فقط: {n_name})")


if __name__ == "__main__":
    main()
