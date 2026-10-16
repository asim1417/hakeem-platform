"""ملف معاينة الحقن للمالك (ف١١: أرقام وعيّنة قبل أي كتابة). لا يكتب على أي قاعدة.
  python3 scripts/build_injection_xlsx.py <scratch_dir> <injection_dir> <out.xlsx>
"""
import collections, json, pathlib, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

S, D, OUT = pathlib.Path(sys.argv[1]), pathlib.Path(sys.argv[2]), sys.argv[3]
jl = lambda f: [json.loads(l) for l in open(D / f, encoding='utf-8')]
F, FB = Font(name='Arial', size=10), Font(name='Arial', size=10, bold=True)
YEL, GREY = PatternFill('solid', fgColor='FFF2CC'), PatternFill('solid', fgColor='E7E6E6')
CLIP = 32000
wb = Workbook()


def sheet(title, head, rows, widths, decision=None):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.append(head)
    for r in rows:
        ws.append([(str(v)[:CLIP] if isinstance(v, str) else v) for v in r])
    for c in ws[1]:
        c.font, c.fill = FB, GREY
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment = F, Alignment(wrap_text=True, vertical='top')
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    if decision:
        dv = DataValidation(type='list', formula1=f'"{decision}"', allow_blank=True)
        ws.add_data_validation(dv)
        col = len(head)
        for r in range(2, ws.max_row + 1):
            ws.cell(r, col).fill = YEL
            dv.add(ws.cell(r, col))
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = ws.dimensions
    return ws


summ = json.load(open(D / 'summary.json'))
im = json.load(open(S / 'identity_match.json'))
tc = json.load(open(S / 'text_compare.json'))
dc = json.load(open(S / 'diff_classes.json'))
cur = {r['id']: r['c'] for r in json.load(open(S / 'uv_current.json'))}

ws0 = wb.active
ws0.title = 'ملخص'
ws0.sheet_view.rightToLeft = True
g = collections.Counter(o['grade'] for o in im)
tcc = collections.Counter()
for o in tc:
    tcc.update(o['counts'])
dcc = collections.Counter(r['class'] for r in dc)
rows = [
    ['معاينة حقن نتائج هيئة الخبراء في قاعدة أنظمة حكيم — 16 أكتوبر 2026'],
    ['لم يُكتب شيء في الإنتاج. الحقن ينتظر كلمة المالك «احقن» بعد مراجعة هذا الملف.'],
    [],
    ['أ. الواقع الحالي في حكيم (قراءة من الإنتاج اليوم)', 'العدد'],
    ['أنظمة', 1146], ['أنظمة لها مواد', 1063], ['مواد', 28606], ['مواد حالتها «ملغاة»', 264],
    ['أنظمة لها مقدمة', 990], ['أنظمة بلا تصنيف', 642], ['نسخ مواد مؤرخة (unit_version)', 0],
    ['صفوف تحقق (verification)', 715], ['أنظمة لها حالة متحقق منها', 28],
    ['روابط اللوائح بأنظمتها (IMPLEMENTS)', 0], ['جداول الهيئة (boe) في الإنتاج', 'غير موجودة'],
    ['حارسا md5 (التنفيذ م1، المعاملات المدنية م1)', 'مطابقان'],
    [],
    ['ب. ما جمعناه من هيئة الخبراء (لقطة 28/9/2026)', 'العدد'],
    ['أنظمة', 530], ['مواد', 16859], ['ملاحظات تعديل', 1819], ['اقتراحات مراجعة', 741],
    ['قراءات صور', 166], ['مرفقات مقروءة / صفحاتها', '982 / 8,860'], ['مرفقات لم تُقرأ', 2142],
    [],
    ['ج. مطابقة الهوية (الاسم + أداة الإصدار + التاريخ)', 'العدد'],
    ['مؤكدة (الدرجة ١) — يُحقن فيها', g[1]], ['اسم بلا أداة مطابقة (الدرجة ٤) — للمالك', g[4]],
    ['أكثر من مرشح (الدرجة ٥) — للمالك', g[5]], ['غير موجودة في حكيم', g[None]],
    [],
    ['د. مقارنة نص المواد داخل الأزواج المؤكدة', 'العدد'],
    ['مطابقة حرفيًّا', tcc['مطابقة']], ['فرق شكلي (عنوان المادة، ترقيم البنود، التشكيل)', dcc['شكلي']],
    ['فرق لفظي طفيف', dcc['لفظي طفيف']], ['فرق لفظي', dcc['لفظي']], ['مختلف جوهريًّا', dcc['مختلف جوهريًّا']],
    ['ناقصة في حكيم', tcc['ناقصة في حكيم']], ['زائدة في حكيم', tcc['زائدة في حكيم']],
    [],
    ['هـ. الجاهز للحقن (اجتاز التحقق المستقل)', 'الصفوف'],
    ['حالة النظام «ساري»', summ['حالة النظام'].get('ساري', 0)],
    ['حالة النظام «ملغى» (يظهر اليوم في حكيم كأنه ساري)', summ['حالة النظام'].get('ملغى', 0)],
    ['نسخ مواد معدّلة (نص كامل جديد بتاريخه)', summ['نسخ مواد']], ['مواد تأخذ نسخة', summ['مواد بنسخ']],
    ['إلغاء مواد', summ['إلغاء مواد']],
    ['المجموع', summ['نسخ مواد'] + summ['إلغاء مواد'] + sum(summ['حالة النظام'].values())],
    [],
    ['و. المؤجل وسببه (ملاحظات)', 'العدد'],
] + [[k, v] for k, v in sorted(summ['مؤجل (ملاحظات)'].items(), key=lambda t: -t[1])] + [
    ['تعارض حالة النظام بين حكيم والهيئة', summ['تعارض حالة مع حكيم']],
]
for r in rows:
    ws0.append(r)
ws0.column_dimensions['A'].width = 70
ws0.column_dimensions['B'].width = 18
for row in ws0.iter_rows():
    for c in row:
        c.font = F
for r in ws0.iter_rows():
    if r[0].value and str(r[0].value)[:2] in ('أ.', 'ب.', 'ج.', 'د.', 'هـ', 'و.'):
        for c in r:
            c.font, c.fill = FB, GREY
ws0['A1'].font = Font(name='Arial', size=12, bold=True)

ws = sheet('حالة الأنظمة', ['النظام', 'الحالة المقترحة', 'أداة الإصدار', 'رابط الهيئة', 'قرار المالك'],
           [[r['_law'], r['verified_status'], r['evidence_instrument'], r['evidence_url'], None] for r in
            sorted(jl('work_status.jsonl'), key=lambda r: (r['verified_status'] != 'ملغى', r['_law']))],
           [50, 14, 30, 40, 14], 'موافق,مرفوض')
sheet('نسخ المواد', ['النظام', 'المادة', 'أداة التعديل', 'يسري من (ميلادي)', 'النص الحالي في حكيم', 'النص الجديد', 'ملاحظة الهيئة (الدليل)', 'رابط', 'قرار المالك'],
      [[r['_law'], r['_article'], r['evidence_instrument'], r['valid_from'], cur.get(r['unit_id']), r['body'], r['evidence_quote'], r['evidence_url'], None]
       for r in sorted(jl('unit_versions.jsonl'), key=lambda r: (r['_law'], r['_article'], r['valid_from']))],
      [32, 8, 28, 13, 55, 55, 45, 30, 14], 'موافق,مرفوض,موافق مع تصحيح')
sheet('إلغاء المواد', ['النظام', 'المادة', 'أداة الإلغاء', 'ملاحظة الهيئة', 'ملغاة أصلًا في حكيم', 'رابط', 'قرار المالك'],
      [[r['_law'], r['_article'], r['evidence_instrument'], r['evidence_quote'], 'نعم' if r['_already_repealed_in_hakeem'] else 'لا', r['evidence_url'], None]
       for r in sorted(jl('unit_status.jsonl'), key=lambda r: (r['_law'], r['_article']))],
      [34, 8, 30, 50, 12, 30, 14], 'موافق,مرفوض')
sheet('تعارض الحالة', ['النظام', 'الحالة في حكيم', 'الحالة في الهيئة', 'قرار المالك'],
      [[r['law'], '، '.join(r['hakeem']), r['boe'], None] for r in jl('work_conflicts.jsonl')],
      [50, 30, 20, 20], 'حالة حكيم,حالة الهيئة,مؤجل')
sheet('مطابقة للمالك', ['نظام الهيئة', 'الحالة', 'مواده', 'الدرجة', 'المرشحون في حكيم', 'أداة الهيئة', 'قرار المالك'],
      [[o['boe_name'], o['boe_status'], o['boe_units'],
        {4: 'اسم بلا أداة', 5: 'أكثر من مرشح', None: 'غير موجود'}[o['grade']],
        ' | '.join(c['name'] for c in o['candidates']), ' '.join(map(str, o['boe_instruments'][0])) if o['boe_instruments'] else '', None]
       for o in sorted(im, key=lambda o: (str(o['grade']), o['boe_name'])) if o['grade'] in (4, 5, None)],
      [45, 10, 8, 14, 60, 30, 18], 'ربط بالمرشح الأول,ليس هو,نظام ناقص في حكيم')
fx = collections.defaultdict(collections.Counter)
for r in dc:
    fx[r['law']][r['class']] += 1
sheet('فروق النص', ['النظام', 'مطابقة', 'شكلي', 'لفظي طفيف', 'لفظي', 'مختلف جوهريًّا', 'ناقصة في حكيم', 'زائدة في حكيم'],
      [[o['boe_name'], o['counts'].get('مطابقة', 0), fx[o['boe_name']]['شكلي'], fx[o['boe_name']]['لفظي طفيف'], fx[o['boe_name']]['لفظي'],
        fx[o['boe_name']]['مختلف جوهريًّا'], o['counts'].get('ناقصة في حكيم', 0), o['counts'].get('زائدة في حكيم', 0)]
       for o in sorted(tc, key=lambda o: -(o['counts'].get('مختلفة', 0))) if o['counts'].get('مختلفة') or o['counts'].get('ناقصة في حكيم') or o['counts'].get('زائدة في حكيم')],
      [50, 10, 10, 10, 10, 12, 12, 12])
sheet('المؤجل', ['النظام', 'المادة', 'عدد الملاحظات', 'سبب التأجيل'],
      [[r['law'], r['article'], r['notes'], r['reason']] for r in sorted(jl('deferred.jsonl'), key=lambda r: (r['reason'], r['law'], r['article']))],
      [45, 8, 12, 50])
wb.save(OUT)
print(OUT)
