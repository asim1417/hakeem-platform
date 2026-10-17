"""ورقة قرار المالك: تعارض حالة النظام بين الهيئة وحكيم (8)، وتاريخ نفاذ المواد المؤجلة (30). لا كتابة على أي قاعدة.

  python3 scripts/build_decisions_conflicts_dates.py <prod_work_verification.json>

المدخل الأول: لقطة قراءة من الإنتاج لصفوف verification (work) للأنظمة المتعارضة (حقول object_id، name،
  verified_status، evidence_instrument، evidence_url، quote، verified_by، verified_at).
المخرج: reports/قرارات_التعارض_والنفاذ.xlsx و data/compose-2026-10-16/effective_dates.jsonl

قواعد تاريخ النفاذ:
- يُبحث عن عبارة النفاذ في نص الملاحظة بعد حذف ما بين علامات التنصيص (نص المادة نفسه لا يُحسب).
- «ويعمل بها/بالتعديل/به من تاريخ X» ← valid_from = X بتقويم أم القرى، الأساس explicit.
- «بعد (N) يومًا من تاريخ نشره» ← يحتاج تاريخ النشر في أم القرى (gazette)؛ لا يُقترح تاريخ.
- لا عبارة نفاذ خارج نص المادة ← إنذار كاذب: يبقى تاريخ الأداة (issue).
"""
import json
import pathlib
import re
import sys

from hijridate import Hijri
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = pathlib.Path(__file__).resolve().parents[1]
INJ = ROOT / "data" / "injection-2026-10-16"
COMPOSE = ROOT / "data" / "compose-2026-10-16"
XLSX = ROOT / "reports" / "قرارات_التعارض_والنفاذ.xlsx"

QUOTED = re.compile(r"[\"“«][^\"”»]*[\"”»]")
EXPLICIT = re.compile(r"(?:ويعمل|يعمل|ويُعمل|يُعمل)\s+(?:بها|به|بالتعديل|بالتعديلات|بهذا التعديل)\s+(?:اعتبارًا\s+|اعتبارا\s+)?من\s+تاريخ\s+(\d{4})/(\d{1,2})/(\d{1,2})")
AFTER_PUB = re.compile(r"بعد\s+\(?\s*([^)\s]+)\s*\)?\s+يومًا?\s+من\s+تاريخ\s+نشره")


def jl(p):
    with open(p, encoding="utf-8") as f:
        return [json.loads(x) for x in f if x.strip()]


def to_g(y, m, d):
    try:
        return Hijri(int(y), int(m), int(d)).to_gregorian().isoformat()
    except (ValueError, OverflowError):
        return None


def effective_for(row):
    for n in reversed(row.get("notes", [])):
        outside = QUOTED.sub(" ", n["text"] or "")
        m = EXPLICIT.search(outside)
        if m:
            return {"basis": "explicit", "hijri": f"{m[1]}/{int(m[2])}/{int(m[3])}", "valid_from": to_g(m[1], m[2], m[3]),
                    "clause": m.group(0), "note_id": n["id"]}
        a = AFTER_PUB.search(outside)
        if a:
            return {"basis": "gazette", "hijri": None, "valid_from": None, "clause": a.group(0), "note_id": n["id"],
                    "detail": "يحتاج تاريخ النشر في أم القرى"}
    return {"basis": "issue", "hijri": None, "valid_from": row.get("valid_from"), "clause": None,
            "detail": "إنذار كاذب: عبارة المدة من نص المادة لا من الملاحظة — يبقى تاريخ الأداة"}


def header(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"


def main():
    prod = {r["object_id"]: r for r in json.load(open(sys.argv[1], encoding="utf-8"))}
    conflicts = jl(INJ / "work_conflicts.jsonl")
    cands = [r for r in jl(COMPOSE / "candidates.jsonl") if r["deferred_reason"].startswith("تاريخ النفاذ")]

    wb = Workbook()
    yellow = PatternFill("solid", fgColor="FFF2CC")
    ws = wb.active
    ws.title = "تعارض الحالة"
    header(ws, ["النظام", "حالة الهيئة (28/9/2026)", "حالة حكيم", "سند حكيم", "نص السند", "رابط السند", "التوصية", "السبب",
                "قرار المالك", "ملاحظة"])
    for c in conflicts:
        p = prod.get(c["system_id"], {})
        boe = c["boe"]
        rec = "أبقِ «مستبدل»"
        why = ("أداة الاستبدال صريحة في أم القرى وسريانها مضى؛ الهيئة لم تحدّث الحالة بعد."
               if boe == "ساري" else "«مستبدل» صورة من الإلغاء، والهيئة تقول «لاغي»: لا تعارض في المعنى.")
        ws.append([c["law"], boe, "، ".join(c["hakeem"]), p.get("evidence_instrument"), p.get("quote"), p.get("evidence_url"),
                   rec, why, "", ""])
        ws.cell(ws.max_row, 9).fill = yellow

    ws2 = wb.create_sheet("تاريخ النفاذ")
    header(ws2, ["النظام", "المادة", "حالة التركيب", "تاريخ الأداة (issue)", "الأساس المقترح", "النفاذ الهجري",
                 "يسري من (مقترح)", "عبارة النفاذ في الملاحظة", "تفصيل", "قرار المالك", "تاريخ يحدده المالك"])
    out = []
    for r in sorted(cands, key=lambda r: (r["law"], r["article"])):
        e = effective_for(r)
        out.append({"system_id": r["system_id"], "law": r["law"], "article": r["article"], "candidate_id": r.get("id"),
                    "issue_valid_from": r.get("valid_from"), **e})
        ws2.append([r["law"], r["article"], r["compose_status"], r.get("valid_from"), e["basis"], e["hijri"], e["valid_from"],
                    e["clause"], e.get("detail"), "", ""])
        ws2.cell(ws2.max_row, 10).fill = yellow
        ws2.cell(ws2.max_row, 11).fill = yellow
    for sheet in (ws, ws2):
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 28
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX)
    with open(COMPOSE / "effective_dates.jsonl", "w", encoding="utf-8") as f:
        for r in out:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    from collections import Counter
    print(json.dumps({"تعارضات": len(conflicts), "تاريخ النفاذ": dict(Counter(r["basis"] for r in out))}, ensure_ascii=False))


if __name__ == "__main__":
    main()
