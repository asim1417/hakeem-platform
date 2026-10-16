"""تقرير المالك من قاعدة الهيئة نفسها (بلا حاجة لقاعدة حكيم): الأنظمة، والتعديلات، والملغى، والشاذ،
وتغطية قائمة حكيم المرجعية (500 اسم) كتقدير أولي. الربط الملزم يُحسب على قاعدة حكيم بـ match_and_diff.py."""
import collections
import csv
import json
import pathlib
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(pathlib.Path(__file__).parent))
from common import normalize_system_name  # noqa: E402

ROOT = pathlib.Path(__file__).resolve().parents[1]
P = ROOT / "data" / "parsed"
EFFECT_AR = {
    "amend_article": "تعديل مادة", "amend_paragraph": "تعديل فقرة", "amend_definition": "تعديل تعريف", "amend_table": "تعديل جدول",
    "renumber_merge": "دمج أو إعادة ترتيب", "repeal": "إلغاء", "delete": "حذف", "add": "إضافة", "substitute_phrase": "إحلال عبارة",
    "general": "ملاحظة عامة", "unclassified": "غير مصنف",
}
KIND_AR = {"royal_decree": "مرسوم ملكي", "cabinet_decision": "قرار مجلس الوزراء", "royal_order": "أمر ملكي"}


def jl(name):
    return [json.loads(x) for x in open(P / f"{name}.jsonl", encoding="utf-8") if x.strip()]


def main(ref_csv: str | None) -> None:
    laws = jl("law")
    units = jl("unit")
    notes = jl("amendment_note")
    rep = json.loads((P / "report.json").read_text(encoding="utf-8"))
    catalog = {c["guid"]: c for c in json.loads((ROOT / "data" / "boe_catalog.json").read_text(encoding="utf-8"))}
    by_law_units = collections.defaultdict(list)
    for u in units:
        by_law_units[u["law_id"]].append(u)
    by_law_notes = collections.defaultdict(list)
    for n in notes:
        by_law_notes[n["law_id"]].append(n)
    wb = Workbook()
    head = Font(name="Arial", bold=True)
    body = Font(name="Arial")
    yellow = PatternFill("solid", fgColor="FFF2CC")

    def sheet(title, headers, rows, widths=None, owner=()):
        ws = wb.create_sheet(title)
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for c in ws[1]:
            c.font = head
        for r in rows:
            ws.append(r)
        for i, h in enumerate(headers, start=1):
            col = ws.cell(row=1, column=i).column_letter
            ws.column_dimensions[col].width = (widths or {}).get(h, 22)
            if h in owner:
                for row in ws.iter_rows(min_col=i, max_col=i):
                    for c in row:
                        c.fill = yellow
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = body
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    # الأنظمة
    law_rows = []
    for l in sorted(laws, key=lambda x: (catalog.get(x["guid"], {}).get("tab", ""), catalog.get(x["guid"], {}).get("category", ""), x["name"] or "")):
        us = by_law_units[l["id"]]
        ns = by_law_notes[l["id"]]
        c = catalog.get(l["guid"], {})
        law_rows.append([
            c.get("tab"), " › ".join(l["category_path"]), l["name"], l["status_at_source"], l["issued_hijri"], l["issued_gregorian"],
            l["published_hijri"], "؛ ".join(i["text"] for i in l["instruments"]), len(us),
            sum(1 for u in us if u["css_class"] == "changed-article"), sum(1 for u in us if u["css_class"] == "canceled-article"),
            len(ns), sum(1 for n in ns if n["needs_review"]), "، ".join(sorted({f.split(":")[0] for f in l["parse_flags"]})), (l["summary"] or "")[:1500], l["url"],
        ])
    for g in rep.get("not_found", []):
        c = catalog.get(g, {})
        law_rows.append([c.get("tab"), c.get("category"), c.get("name"), "غير موجود في المصدر", None, None, None, None, 0, 0, 0, 0, 0, "source_not_found", None,
                         f"https://laws.boe.gov.sa/BoeLaws/Laws/LawDetails/{g}/1"])
    # ملخص
    ws0 = wb.active
    ws0.title = "ملخص"
    ws0.sheet_view.rightToLeft = True
    status = collections.Counter(l["status_at_source"] or "—" for l in laws)
    tabs = collections.Counter(catalog.get(l["guid"], {}).get("tab", "—") for l in laws)
    eff = collections.Counter(EFFECT_AR.get(n["effect"], n["effect"]) for n in notes)
    lines = [["البند", "العدد"], ["وثائق الفهرس (الوثائق النظامية)", len(catalog)], ["صفحات مقروءة ومحللة", len(laws)],
             ["صفحات «غير موجود في المصدر»", len(rep.get("not_found", []))], ["مواد", len(units)],
             ["مواد صنفها معدلة (changed-article)", rep["classes"].get("changed-article", 0)],
             ["مواد صنفها ملغاة (canceled-article)", rep["classes"].get("canceled-article", 0)],
             ["ملاحظات تعديل", len(notes)], ["منها تحتاج مراجعة يدوية", sum(1 for n in notes if n["needs_review"])], ["", ""], ["حسب التبويب", ""]]
    lines += [[k, v] for k, v in tabs.most_common()] + [["", ""], ["الحالة في الهيئة", ""]] + [[k, v] for k, v in status.most_common()]
    lines += [["", ""], ["نوع أثر الملاحظة", ""]] + [[k, v] for k, v in eff.most_common()]
    lines += [["", ""], ["تنبيه: الحالة المعروضة في حكيم لا تُكتب من هذا الملف مباشرة؛ تمر بقرار المالك.", ""],
              ["النص المعروض في الهيئة هو الأصل عند الإصدار؛ التعديلات في ورقة «التعديلات».", ""]]
    for r in lines:
        ws0.append(r)
    ws0.column_dimensions["A"].width = 70
    ws0.column_dimensions["B"].width = 14
    for c in ws0[1]:
        c.font = head
    sheet("الأنظمة", ["التبويب", "التصنيف", "الاسم", "الحالة في الهيئة", "الإصدار (هـ)", "الإصدار (م)", "النشر (هـ)", "أدوات الإصدار", "المواد",
                      "معدلة", "ملغاة", "ملاحظات", "تحتاج مراجعة", "وسوم", "نبذة", "الرابط"], law_rows,
          {"الاسم": 45, "التصنيف": 40, "أدوات الإصدار": 45, "نبذة": 60, "الرابط": 40})
    name_of = {l["id"]: l["name"] for l in laws}
    url_of = {l["id"]: l["url"] for l in laws}
    unit_of = {(u["law_id"], u["seq"]): u for u in units}
    note_rows = []
    for n in notes:
        u = unit_of.get((n["law_id"], n["unit_seq"]), {})
        note_rows.append([name_of.get(n["law_id"]), u.get("heading"), u.get("css_class"), EFFECT_AR.get(n["effect"], n["effect"]),
                          KIND_AR.get(n["instrument_kind"] or "", n["instrument_kind"]), n["instrument_no"], n["instrument_date_hijri"],
                          "نعم" if n["needs_review"] else "", n["note_text"][:3000], (n["new_text"] or "")[:3000], "", "", url_of.get(n["law_id"])])
    sheet("التعديلات", ["النظام", "المادة", "صنف المادة", "نوع الأثر", "الأداة", "رقمها", "تاريخها (هـ)", "تحتاج مراجعة", "الملاحظة", "النص الجديد",
                        "قرارك: تُطبَّق/تُرفض/تُؤجَّل", "ملاحظتك", "الرابط"], note_rows,
          {"النظام": 40, "المادة": 22, "الملاحظة": 70, "النص الجديد": 70, "الرابط": 40}, owner=("قرارك: تُطبَّق/تُرفض/تُؤجَّل", "ملاحظتك"))
    retired = [[l["name"], l["status_at_source"], " › ".join(l["category_path"]), "؛ ".join(i["text"] for i in l["instruments"]),
                (l["summary"] or "")[:1500], l["url"], "", "", ""] for l in laws if (l["status_at_source"] or "") != "ساري"]
    sheet("غير الساري في الهيئة", ["الاسم", "الحالة في الهيئة", "التصنيف", "أدوات الإصدار", "نبذة", "الرابط", "حالتك المعتمدة", "الأداة الملغية ورابطها", "النظام الخلف"],
          retired, {"الاسم": 45, "نبذة": 60, "الرابط": 40}, owner=("حالتك المعتمدة", "الأداة الملغية ورابطها", "النظام الخلف"))
    odd = []
    for l in laws:
        us = by_law_units[l["id"]]
        nums = [u["article_number"] for u in us if u["article_number"] and not u["is_bis"]]
        why = [f for f in l["parse_flags"] if not f.startswith("unnumbered")]
        un = [u["heading"] for u in us if u["article_number"] is None]
        if un:
            why.append(f"وحدات بلا رقم مادة: {len(un)} ({'، '.join(un[:4])}{'…' if len(un) > 4 else ''})")
        if nums != sorted(nums):
            why.append("ترقيم غير متصاعد")
        if len(nums) != len(set(nums)):
            why.append("رقم مادة مكرر")
        bis = [u["heading"] for u in us if u["is_bis"]]
        if bis:
            why.append(f"مواد مكرر: {'، '.join(bis[:5])}")
        if why:
            odd.append([l["name"], "؛ ".join(why), l["url"]])
    for g in rep.get("not_found", []):
        odd.append([catalog.get(g, {}).get("name"), "الصفحة ردّت «العنصر المطلوب غير موجود بالنظام»", f"https://laws.boe.gov.sa/BoeLaws/Laws/LawDetails/{g}/1"])
    sheet("شاذ يحتاج مراجعة", ["النظام", "السبب", "الرابط"], odd, {"النظام": 45, "السبب": 80, "الرابط": 40})
    if ref_csv and pathlib.Path(ref_csv).exists():
        boe_by = collections.defaultdict(list)
        for l in laws:
            boe_by[normalize_system_name(l["name"] or "")].append(l)
        rows = []
        with open(ref_csv, encoding="utf-8") as f:
            for r in csv.DictReader(f):
                nm = r.get("اسم النظام / الأداة") or ""
                hits = boe_by.get(normalize_system_name(nm), [])
                rows.append([r.get("الرمز"), nm, r.get("النوع"), len(hits), "؛ ".join(h["name"] for h in hits), "؛ ".join(h["status_at_source"] or "" for h in hits),
                             "؛ ".join(h["url"] for h in hits)])
        sheet("تغطية قائمة حكيم (تقدير)", ["الرمز", "اسم حكيم", "النوع", "مطابقات الهيئة", "اسم الهيئة", "الحالة في الهيئة", "الرابط"], rows,
              {"اسم حكيم": 45, "اسم الهيئة": 45, "الرابط": 40})
        ws0.append(["", ""])
        ws0.append(["قائمة حكيم المرجعية (500): أسماء لها مطابق واحد بالاسم المطبَّع في الهيئة", sum(1 for r in rows if r[3] == 1)])
        ws0.append(["قائمة حكيم المرجعية: بلا مطابق بالاسم (قد تكون لائحة أو تسمية مختلفة أو غير موجودة)", sum(1 for r in rows if r[3] == 0)])
    out = ROOT / "reports" / "قاعدة_هيئة_الخبراء_تقرير_المالك.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    print(out, len(law_rows), len(note_rows), len(retired), len(odd))


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else None)
