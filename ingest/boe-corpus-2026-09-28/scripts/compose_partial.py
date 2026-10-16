"""المرحلة ٤ — تركيب النص الكامل للمواد المؤجلة (تعديل فقرة أو عبارة). لا شبكة ولا قاعدة.

  python3 scripts/compose_partial.py

المدخلات: data/parsed/{law,unit,amendment_note}.jsonl (من parse_snapshots.py)،
          data/review/proposals.jsonl، data/injection-2026-10-16/{deferred,work_status}.jsonl
المخرج:  data/compose-2026-10-16/candidates.jsonl و summary.json
         reports/تركيب_المواد.xlsx (فرق قبل وبعد، وعمود قرار أصفر)

القواعد (أمر الحقن §المرحلة ٤):
- الأساس نص الهيئة الأصلي (unit.body)، وتُطبَّق التعديلات بتاريخ الأداة من الأقدم.
- إعادة نص المادة كاملًا داخل الملاحظة («ويكون نص المادة… كما يلي») تحل محل المادة.
- تعديل مادة كاملة بنص جديد يحل محلها.
- إحلال عبارة أو حذفها: بشرط ورود القديمة مرة واحدة بالضبط، أو «أينما/حيثما وردت».
- تعديل فقرة برمزها: يُستبدل ما بين بدايتها وبداية الفقرة التالية، بشرط موضع واحد.
- ما سوى ذلك، أو موضع غير موجود أو متعدد: compose_status = manual مع السبب، ولا يُكتب نص.
- التطابق على نص مطبَّع، والاستبدال على النص الأصلي بمواضعه. لا نص من OCR.
"""
import collections
import difflib
import hashlib
import json
import pathlib
import re

from hijridate import Hijri

ROOT = pathlib.Path(__file__).resolve().parents[1]
PARSED = ROOT / "data" / "parsed"
INJ = ROOT / "data" / "injection-2026-10-16"
OUT = ROOT / "data" / "compose-2026-10-16"
XLSX = ROOT / "reports" / "تركيب_المواد.xlsx"
RESOLVED = {"محسوم من الملاحظة", "مؤكد"}
NO_TEXT = {"إجراء إداري لا يمس النص", "تفسير", "general"}

DIAC = re.compile(r"[ً-ْٰـ]")
QUOTE = r"[\"«“(\[]"
UNQUOTE = r"[\"»”)\]]"


def sid(*parts):
    return "boe26-" + hashlib.sha256("|".join(map(str, parts)).encode()).hexdigest()[:24]


def jl(p):
    with open(p, encoding="utf-8") as f:
        return [json.loads(x) for x in f if x.strip()]


def norm_char(c: str) -> str:
    return {"أ": "ا", "إ": "ا", "آ": "ا", "ى": "ي", "ة": "ه"}.get(c, c)


def normalized_with_map(text: str) -> tuple[str, list[int]]:
    """نص مطبَّع وخريطة من كل حرف فيه إلى موضعه في الأصل (لاستبدال على الأصل بمواضعه)."""
    out, idx = [], []
    prev_space = False
    for i, c in enumerate(text):
        if DIAC.match(c):
            continue
        if c.isspace():
            if prev_space:
                continue
            out.append(" ")
            idx.append(i)
            prev_space = True
            continue
        prev_space = False
        out.append(norm_char(c))
        idx.append(i)
    return "".join(out), idx


def norm(s: str) -> str:
    return normalized_with_map(s)[0].strip()


def find_all(text: str, needle: str) -> list[tuple[int, int]]:
    """مواضع needle في text على المطبَّع، مُعادة بمواضع الأصل [start, end)."""
    nt, idx = normalized_with_map(text)
    nn = norm(needle)
    if not nn:
        return []
    spans, pos = [], 0
    while True:
        k = nt.find(nn, pos)
        if k < 0:
            break
        start = idx[k]
        end = idx[k + len(nn) - 1] + 1
        spans.append((start, end))
        pos = k + len(nn)
    return spans


def greg(h):
    nums = re.findall(r"\d+", h or "")
    if len(nums) < 3:
        return None
    a, b, c = map(int, nums[:3])
    d, m, y = (c, b, a) if a > 1000 else (a, b, c)
    try:
        return Hijri(y, m, d).to_gregorian().isoformat()
    except (ValueError, OverflowError):
        return None


def strip_tail(t: str) -> str:
    t = t.strip()
    t = re.split(r"\n\s*(?:وذلك بموجب|-\s*تم |تم تعديل|تم إلغاء|عُدّلت|عدلت)", t)[0]
    t = re.split(r"[\"»”]\s*[.،]?\s*(?=وذلك بموجب|بموجب (?:المرسوم|الأمر|قرار))", t)[0]
    t = t.strip().strip("\"«»“” \n")
    # علامة إغلاق قبل نقطة أو فاصلة ختامية («…الإنتاجية".») تُزال وتبقى العلامة.
    t = re.sub(r"[\"»”]+\s*([.،؛]?)\s*$", r"\1", t)
    return t.strip()


def unbalanced_quote(note_text: str) -> bool:
    """علامة تنصيص مفتوحة بلا إغلاق: النص قد يكون مقتطعًا."""
    t = note_text or ""
    return t.count("«") != t.count("»") or t.count("“") != t.count("”") or t.count('"') % 2 == 1


def clean_body(t: str) -> str:
    t = (t or "").strip().strip("\"«»“” \n")
    t = re.sub(r"^\s*المادة\s+[^\n:]{1,60}:\s*", "", t).strip()
    return strip_tail(t)


# ── أنماط الملاحظات ──
FULL_RESTATE = re.compile(
    r"(?:ويكون نص المادة|ليكون نص المادة|ليصبح نص المادة|لتصبح المادة|وإحلال النص الآتي محله|وإحلال النص التالي محله"
    r"|لتكون بالنص (?:الآتي|التالي)|ليكون بالنص (?:الآتي|التالي)(?= *:?\s*\n*\s*[\"«]?\s*المادة))[^:\n]{0,40}:\s*",
)
SUB_NEW_FOR_OLD = re.compile(
    rf"(?:إحلال|حلت|تحل|يحل|وإحلال|و)\s*عبارة\s*{QUOTE}\s*(?P<new>[^\"»”)\]]{{1,200}}?)\s*{UNQUOTE}\s*(?:محل|مكان|بدلًا من|بدلا من)\s*(?:عبارة\s*)?{QUOTE}\s*(?P<old>[^\"»”)\]]{{1,200}}?)\s*{UNQUOTE}"
)
SUB_OLD_WITH_NEW = re.compile(
    rf"استبدال\s*(?:عبارة\s*)?{QUOTE}\s*(?P<old>[^\"»”)\]]{{1,200}}?)\s*{UNQUOTE}\s*(?:بعبارة|ب)\s*{QUOTE}\s*(?P<new>[^\"»”)\]]{{1,200}}?)\s*{UNQUOTE}"
)
DELETE_PHRASE = re.compile(rf"حذف\s*(?:عبارة|كلمة)\s*{QUOTE}\s*(?P<old>[^\"»”)\]]{{1,200}}?)\s*{UNQUOTE}")
AMEND_PHRASE = re.compile(
    rf"(?:بتعديل|تعديل)\s+(?:عبارة|كلمة)\s*{QUOTE}\s*(?P<old>[^\"»”)\]]{{1,200}}?)\s*{UNQUOTE}\s*(?:لتكون|لتصبح|ليكون|إلى)\s*{QUOTE}\s*(?P<new>[^\"»”)\]]{{1,200}}?)\s*{UNQUOTE}"
)
INSERT_AFTER = re.compile(
    rf"إضافة\s+(?:عبارة|كلمة|العبارة)\s*{QUOTE}\s*(?P<new>[^\"»”)\]]{{1,300}}?)\s*{UNQUOTE}\s*بعد\s+(?:عبارة|كلمة)\s*{QUOTE}\s*(?P<anchor>[^\"»”)\]]{{1,200}}?)\s*{UNQUOTE}"
)
# أفعال العمليات في الملاحظة (لا تُحسب صيغة الخبر «عُدلت/تم تعديل»)
OP_VERB = re.compile(r"(?<![\u0621-\u064A])(?:بتعديل|تعديل|إضافة|حذف|إلغاء|إحلال|استبدال)\s+(?!هذه المادة|المادة|نص هذه|الفقرة|البند)")
EVERYWHERE = re.compile(r"(?:أينما|حيثما|اينما|حيثُما)\s+وردت")
PARA_AMEND = re.compile(
    r"(?:تعديل|عُدلت|عدلت)\s+(?:نص\s+)?(?:الفقرة|البند)\s*\(\s*(?P<label>[^)]{1,8}?)\s*\)[^\n]{0,80}?(?:(?:ليكون|لتكون|لتصبح|ليصبح)\s+(?:نصها\s+|نصه\s+)?(?:بالنص\s+)?(?:التالي|الآتي|كما يلي)|على النحو (?:الآتي|التالي))\s*:?\s*(?P<body>.+)",
    re.S,
)
LABEL_LINE = re.compile(r"^\s*\(?\s*(?:[أ-ي]{1,2}|هـ|\d{1,2}|أولاً|ثانياً|ثالثاً|رابعاً|خامساً|سادساً|سابعاً|ثامناً|تاسعاً|عاشراً)\s*\)?\s*[-–ـ/):.]", re.M)


LEAK = re.compile(r"وذلك بموجب|بموجب (?:المرسوم|الأمر|قرار)|ويعمل بهذا|يُعمل بهذا|عُدِّلت|عدلت هذه|تم تعديل")
# ما يدل على أن التعديل جزئي (من build_injection.py): لا يُستبدل به نص المادة كاملًا.
PARTIAL_NOTE = re.compile(r"صدر (هذه )?المادة|عجز (هذه )?المادة|فقر|البند|بند|عبارة|إحلال|الشطر")
WHOLE_RESTATE = re.compile(r"ويكون نص المادة|ليكون نص المادة|ليصبح نص المادة")


def apply_note(text: str, note: dict) -> tuple[str | None, str]:
    """يطبّق ملاحظة على نص. يعيد (النص الجديد، الوصف) أو (None، سبب الإحالة اليدوية)."""
    nt = note["note_text"] or ""
    effect = note["effect"]

    partial = bool(PARTIAL_NOTE.search(nt))
    if unbalanced_quote(nt):
        return None, "علامة تنصيص مفتوحة بلا إغلاق في الملاحظة (قد يكون النص مقتطعًا)"
    m = FULL_RESTATE.search(nt)
    if m and (not partial or WHOLE_RESTATE.search(nt[m.start():m.end()])):
        body = clean_body(nt[m.end():])
        if len(body) >= 40 and "..." not in body and "…" not in body:
            return body, "إعادة نص المادة كاملًا من الملاحظة"

    if effect in ("amend_article", "تعديل مادة") and note.get("new_text") and not partial:
        body = clean_body(note["new_text"])
        # النص الجديد يجب أن يرد في الملاحظة نفسها؛ ما جاء من مرفق أو قراءة يُعرض على المالك ولا يُركَّب.
        if norm(body) not in norm(nt):
            return None, "النص الجديد ليس في الملاحظة (من مرفق أو قراءة) — يعرض على المالك"
        if len(body) >= 40 and "..." not in body and "…" not in body:
            return body, "تعديل مادة كاملة"
        return None, "نص المادة الجديد مقتطع أو قصير"

    pairs = [(x["old"], x["new"]) for x in SUB_NEW_FOR_OLD.finditer(nt)]
    pairs += [(x["old"], x["new"]) for x in SUB_OLD_WITH_NEW.finditer(nt)]
    pairs += [(x["old"], x["new"]) for x in AMEND_PHRASE.finditer(nt)]
    inserts = [(x["anchor"], x["new"]) for x in INSERT_AFTER.finditer(nt)]
    ops = len(pairs) + len(inserts)
    if ops and len(OP_VERB.findall(nt)) > ops:
        return None, "ملاحظة بعمليات متعددة لم تُفهم كلها"
    if inserts and not pairs:
        out = text
        for anchor, new in inserts:
            spans = find_all(out, anchor)
            if len(spans) != 1:
                return None, f"موضع الإضافة ورد {len(spans)} مرة: «{anchor}»"
            e = spans[0][1]
            out = out[:e] + " " + new.strip() + out[e:]
        return out, "إضافة عبارة: " + "؛ ".join(f"بعد «{a}» ← «{n}»" for a, n in inserts)
    if pairs and inserts:
        return None, "ملاحظة تجمع إحلالًا وإضافة"
    if pairs:
        every = bool(EVERYWHERE.search(nt))
        out = text
        for old, new in pairs:
            spans = find_all(out, old)
            if not spans:
                return None, f"العبارة القديمة غير موجودة: «{old}»"
            if len(spans) > 1 and not every:
                return None, f"العبارة القديمة وردت {len(spans)} مرات بلا «أينما وردت»: «{old}»"
            for s, e in reversed(spans):
                out = out[:s] + new.strip() + out[e:]
        return out, "إحلال عبارة: " + "؛ ".join(f"«{o}» ← «{n}»" for o, n in pairs)

    d = DELETE_PHRASE.search(nt)
    if d and effect in ("delete", "substitute_phrase", "unclassified"):
        spans = find_all(text, d["old"])
        if len(spans) != 1 and not (spans and EVERYWHERE.search(nt)):
            return None, f"العبارة المحذوفة وردت {len(spans)} مرة: «{d['old']}»"
        out = text
        for s, e in reversed(spans):
            out = out[:s] + out[e:]
        return re.sub(r"[ \t]{2,}", " ", out), f"حذف عبارة «{d['old']}»"

    p = PARA_AMEND.search(nt)
    if p and (OP_VERB.search(nt) or len(re.findall(r"(?:الفقرة|الفقرات|البند|البنود)\s*\(", nt)) > 1):
        return None, "ملاحظة فقرة بعمليات أو فقرات متعددة"
    if p:
        label = p["label"].strip()
        body = strip_tail(p["body"])
        if "ويكون نص المادة" in body:
            body = body.split("ويكون نص المادة")[0].strip()
        body = body.strip("\"«»“”() \n.")
        starts = [mm for mm in re.finditer(rf"^\s*\(?\s*{re.escape(label)}\s*\)?\s*[-–ـ/):.]", text, re.M)]
        if len(starts) != 1:
            return None, f"الفقرة ({label}) وُجدت {len(starts)} مرة"
        s = starts[0].start()
        nxt = [mm.start() for mm in LABEL_LINE.finditer(text) if mm.start() > starts[0].end()]
        e = nxt[0] if nxt else len(text)
        lead = text[s:starts[0].end()]
        new_para = body if norm(body).startswith(norm(label)) else f"{lead.strip()} {body}"
        seg = text[s:e]
        tail = seg[len(seg.rstrip()):] if e < len(text) else ""
        return text[:s] + new_para.strip() + (tail or ("\n" if e < len(text) else "")) + text[e:], f"تعديل الفقرة ({label})"

    return None, f"أثر «{effect}» بلا نمط تركيب آلي"


def diff_text(a: str, b: str) -> str:
    aw, bw = a.split(), b.split()
    out = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(a=aw, b=bw, autojunk=False).get_opcodes():
        if op == "equal":
            out.append(" ".join(aw[i1:i2]))
        if op in ("delete", "replace"):
            out.append("[-" + " ".join(aw[i1:i2]) + "-]")
        if op in ("insert", "replace"):
            out.append("[+" + " ".join(bw[j1:j2]) + "+]")
    return " ".join(out)


def main() -> None:
    laws = {r["id"]: r for r in jl(PARSED / "law.jsonl")}
    guid_to_law = {r["guid"]: r for r in laws.values()}
    name_to_law = collections.defaultdict(list)
    for r in laws.values():
        name_to_law[norm(r["name"])].append(r)
    units, seen = {}, collections.Counter()
    for u in jl(PARSED / "unit.jsonl"):
        if u["article_number"] is not None and not u["is_bis"]:
            seen[(u["law_id"], u["article_number"])] += 1
            units[(u["law_id"], u["article_number"])] = u
    # رقم مكرر في صفحة الهيئة = أكثر من إصدار للنظام في الصفحة نفسها (مثل باب العقوبات في نظام العمل):
    # لا يُعرف أي النصين أصل المادة الحالية، فلا تركيب آلي.
    dup_numbers = {k for k, v in seen.items() if v > 1}
    notes_by_unit = collections.defaultdict(list)
    for n in jl(PARSED / "amendment_note.jsonl"):
        notes_by_unit[(n["law_id"], n["unit_seq"])].append(n)
    proposals = {}
    for p in jl(ROOT / "data" / "review" / "proposals.jsonl"):
        proposals.setdefault(p["note_id"].partition("#")[0], p)

    sys_guid = {}
    for w in jl(INJ / "work_status.jsonl"):
        m = re.search(r"LawDetails/([0-9a-f-]{36})", w.get("evidence_url") or "")
        if m:
            sys_guid[w["object_id"]] = m[1]

    deferred = jl(INJ / "deferred.jsonl")
    cands, stats = [], collections.Counter()
    for d in deferred:
        law = guid_to_law.get(sys_guid.get(d["system_id"], ""))
        if not law:
            same = name_to_law.get(norm(d["law"]), [])
            law = same[0] if len(same) == 1 else None
        row = {"system_id": d["system_id"], "law": d["law"], "article": d["article"], "deferred_reason": d["reason"]}
        if not law:
            row.update(compose_status="manual", compose_detail="النظام غير مربوط بصفحة الهيئة")
            cands.append(row)
            stats["manual: النظام غير مربوط"] += 1
            continue
        if (law["id"], d["article"]) in dup_numbers:
            row.update(compose_status="manual", compose_detail="رقم المادة مكرر في صفحة الهيئة (أكثر من إصدار)")
            cands.append(row)
            stats["manual: رقم مكرر في الهيئة"] += 1
            continue
        unit = units.get((law["id"], d["article"]))
        if not unit:
            row.update(compose_status="manual", compose_detail="المادة غير موجودة في نص الهيئة")
            cands.append(row)
            stats["manual: المادة غير موجودة في الهيئة"] += 1
            continue
        notes = [n for n in notes_by_unit[(law["id"], unit["seq"])]]
        enriched = []
        for n in notes:
            p = proposals.get(n["id"])
            eff = (p or {}).get("effect") or n["effect"]
            if eff in NO_TEXT or n["effect"] in NO_TEXT:
                continue
            conf = "محسوم من الملاحظة" if not n["needs_review"] else (p or {}).get("confidence") or "بلا اقتراح"
            new_text = (p or {}).get("new_text") or n["new_text"]
            enriched.append({**n, "effect": n["effect"], "new_text": new_text, "confidence": conf,
                             "date_g": greg(n["instrument_date_hijri"])})
        enriched.sort(key=lambda n: (n["date_g"] or "9999", n["note_seq"]))
        text, steps, status, detail = unit["body"], [], "auto", []
        for n in enriched:
            if n["effect"] in ("repeal", "إلغاء مادة") and not re.search(r"فقر|بند", n["note_text"] or ""):
                status, text = "repeal", None
                detail.append("إلغاء المادة (صف تحقق، لا نص)")
                break
            new, why = apply_note(text, n)
            if new is None:
                status = "manual"
                detail.append(why)
                break
            steps.append({"note_id": n["id"], "date_g": n["date_g"], "how": why, "body": new,
                          "instrument": n.get("instrument_kind"), "instrument_no": n.get("instrument_no"),
                          "instrument_hijri": n["instrument_date_hijri"], "note_text": n["note_text"]})
            detail.append(why)
            text = new
        if status == "auto" and text and LEAK.search(text) and not LEAK.search(unit["body"]):
            status = "manual"
            detail.append("صياغة الملاحظة تسربت إلى النص المركّب")
        if status == "auto" and not steps:
            # كل الملاحظات إجرائية لا تمس النص: لا نسخة جديدة.
            status = "no_change"
            detail.append("كل ملاحظاتها لا تمس النص")
        gates = []
        if any(n["confidence"] not in RESOLVED for n in enriched):
            gates.append("ملاحظة غير محسومة")
        if any(n["date_g"] is None for n in enriched):
            gates.append("تاريخ أداة غير مقروء")
        if len({n["date_g"] for n in enriched}) != len(enriched):
            gates.append("تعديلان بتاريخ واحد")
        if status == "auto" and any(st["how"].startswith(("تعديل مادة كاملة", "إعادة نص المادة")) for st in steps):
            base_labels = len(LABEL_LINE.findall(unit["body"]))
            if base_labels >= 2 and len(LABEL_LINE.findall(text or "")) == 0:
                gates.append("نطاق التعديل مشكوك: النص الجديد بلا فقرات والأصل بفقرات")
        if d["reason"] in ("نص حكيم يخالف نص الهيئة الأصلي", "تاريخ النفاذ غير تاريخ الإصدار (ينتظر تحديد تاريخ النفاذ)",
                           "نص الملاحظة متسرب إلى نص المادة", "المادة غير موجودة في حكيم"):
            gates.append(d["reason"])
        last = enriched[-1] if enriched else None
        row.update(
            boe_law_id=law["id"], boe_url=law["url"], unit_seq=unit["seq"], base_body=unit["body"],
            body=text if status == "auto" else None, compose_status=status, compose_detail="؛ ".join(detail),
            steps=steps, notes=[{"id": n["id"], "text": n["note_text"], "effect": n["effect"], "confidence": n["confidence"],
                                 "instrument_hijri": n["instrument_date_hijri"], "date_g": n["date_g"]} for n in enriched],
            valid_from=last["date_g"] if (status == "auto" and last) else None, valid_from_basis="issue",
            gates=gates, ready=status == "auto" and not gates,
        )
        row["id"] = sid("compose", d["system_id"], d["article"], row["valid_from"] or "")
        cands.append(row)
        stats[status] += 1
        if status == "auto" and not gates:
            stats["auto جاهز بلا بوابات"] += 1

    OUT.mkdir(parents=True, exist_ok=True)
    with open(OUT / "candidates.jsonl", "w", encoding="utf-8") as f:
        for r in cands:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    manual_reasons = collections.Counter(re.sub(r"«[^»]*»", "«…»", r["compose_detail"].split("؛ ")[-1]) for r in cands if r["compose_status"] == "manual")
    gates = collections.Counter(g for r in cands if r["compose_status"] == "auto" for g in r.get("gates", []))
    summary = {"مواد مؤجلة": len(deferred), "حالات": dict(stats), "أسباب اليدوي": dict(manual_reasons.most_common()),
               "بوابات المركّب آليًا": dict(gates)}
    (OUT / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=1), encoding="utf-8")
    write_xlsx(cands)
    print(json.dumps(summary, ensure_ascii=False, indent=1))


def write_xlsx(cands: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "التركيب"
    ws.sheet_view.rightToLeft = True
    head = ["المعرف", "النظام", "المادة", "حالة التركيب", "جاهز", "البوابات", "طريقة التركيب / سبب اليدوي",
            "يسري من (تقديري)", "نص الهيئة الأصلي", "النص بعد التركيب", "الفرق", "الملاحظات", "رابط الهيئة",
            "قرار المالك (اعتمد / ارفض / صحّح)", "النص المصحح"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    order = {"auto": 0, "manual": 1, "repeal": 2, "no_change": 3}
    for r in sorted(cands, key=lambda r: (order.get(r["compose_status"], 3), not r.get("ready"), r["law"], r["article"])):
        notes = "\n\n".join(f"[{n['confidence']} · {n['instrument_hijri']}] {n['text']}" for n in r.get("notes", []))
        ws.append([
            r.get("id"), r["law"], r["article"], r["compose_status"], "نعم" if r.get("ready") else "لا",
            "، ".join(r.get("gates", [])), r.get("compose_detail"), r.get("valid_from"),
            (r.get("base_body") or "")[:32000], (r.get("body") or "")[:32000],
            diff_text(r.get("base_body") or "", r.get("body") or "")[:32000] if r.get("body") else "",
            notes[:32000], r.get("boe_url"), "", "",
        ])
        ws.cell(ws.max_row, 14).fill = yellow
        ws.cell(ws.max_row, 15).fill = yellow
    widths = [16, 34, 7, 10, 6, 28, 40, 12, 60, 60, 60, 60, 30, 18, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX)


if __name__ == "__main__":
    main()
