"""
export_issuance_excel.py — الأنظمة وتواريخ إصدارها وأدواتها من بوابة هيئة الخبراء (رصد 28/9/2026).

المصادر (بلا اختلاق — ما لم يرد في المصدر يُترك فارغًا):
  • data/parsed/law.jsonl  ← 530 وثيقة: الاسم، الفئة، تاريخ الإصدار والنشر، أداة الإصدار، الحالة، عدد المواد، الرابط
  • data/injection-2026-10-16/work_status.jsonl ← 416 مطابقة هوية مُتحقَّقة (الهيئة ↔ معرّف النظام في حكيم)
  • data/hakeem_systems_2026-10-16.json ← أسماء أنظمة حكيم (لقطة قراءة من الإنتاج)
  • data/saudi-systems-issuance-2026-09-25.xlsx (الإصدار السابق) ← أنظمة حكيم التي لا مقابل لها في الهيئة

المخرج: reports/saudi-systems-issuance.xlsx (يحلّ محلّ الإصدار السابق)
التشغيل: python3 ingest/boe-corpus-2026-09-28/scripts/export_issuance_excel.py  (يتطلب openpyxl؛
  data/parsed/ غير مُودَع — يُولَّد أولًا بـ scripts/parse_snapshots.py)
"""
import json
import re
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent.parent
ROOT = BASE.parent.parent
DATA = BASE / "data"
OUT = ROOT / "reports" / "saudi-systems-issuance.xlsx"
# الإصدار السابق (485 نظامًا من قاعدة حكيم، 25/9) — مصدر ثابت لا يُكتب فوقه.
PREV = DATA / "saudi-systems-issuance-2026-09-25.xlsx"

KIND = {"royal_decree": "مرسوم ملكي", "cabinet_decision": "قرار مجلس الوزراء", "royal_order": "أمر ملكي", "other": "أخرى"}
STATUS = {"لاغي": "ملغى"}


def norm(s: str) -> str:
    s = re.sub(r"[ً-ْٰـ]", "", s or "")
    s = re.sub(r"[إأآ]", "ا", s).replace("ى", "ي").replace("ة", "ه")
    s = re.sub(r"[()\[\]«»\"'.,،:؛\-–]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def fmt_hijri(h: str | None) -> str:
    """«1423/01/01» ← «1/1/1423هـ»."""
    m = re.match(r"^(\d{4})/(\d{1,2})/(\d{1,2})$", (h or "").strip())
    return f"{int(m[3])}/{int(m[2])}/{m[1]}هـ" if m else ""


def to_date(g: str | None) -> date | None:
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})$", g or "")
    return date(int(m[1]), int(m[2]), int(m[3])) if m else None


def instrument_date(text: str) -> str:
    """تاريخ الأداة كما ورد في نصّها (يوم/شهر/سنة أو سنة/شهر/يوم) بصيغة «ي/ش/سهـ»."""
    t = re.sub(r"\s*/\s*", "/", text)
    m = re.search(r"(\d{1,2})/(\d{1,2})/(1[34]\d\d)", t)
    if m:
        return f"{int(m[1])}/{int(m[2])}/{m[3]}هـ"
    m = re.search(r"(1[34]\d\d)/(\d{1,2})/(\d{1,2})", t)
    return f"{int(m[3])}/{int(m[2])}/{m[1]}هـ" if m else ""


def load_jsonl(p: Path) -> list[dict]:
    return [json.loads(line) for line in p.open(encoding="utf-8") if line.strip()]


def main() -> None:
    laws = load_jsonl(DATA / "parsed" / "law.jsonl")
    works = load_jsonl(DATA / "injection-2026-10-16" / "work_status.jsonl")
    hakeem_name = {r["id"]: r["name"] for r in json.load(open(DATA / "hakeem_systems_2026-10-16.json", encoding="utf-8"))}
    by_guid: dict[str, str] = {}
    for w in works:
        m = re.search(r"LawDetails/([0-9a-f-]+)", w["evidence_url"])
        if m:
            by_guid[m[1]] = w["object_id"]

    # ── الورقة 1: وثائق الهيئة ──
    rows = []
    for law in laws:
        insts = law.get("instruments") or []
        first = insts[0] if insts else None
        hid = by_guid.get(law["guid"])
        cat = law.get("category_path") or []
        issued_h = fmt_hijri(law.get("issued_hijri"))
        inst_h = instrument_date(first["text"]) if first else ""
        year_only = (law.get("issued_hijri") or "").endswith("/01/01") and inst_h != issued_h
        if year_only:
            date_note = "الهيئة تسجّل السنة فقط (1/1)؛ التاريخ الكامل من الأداة" if inst_h else "الهيئة تسجّل السنة فقط (1/1)"
        elif inst_h and inst_h != issued_h:
            date_note = "تاريخ الأداة يختلف عن تاريخ الإصدار المسجّل"
        else:
            date_note = ""
        rows.append({
            "name": law["name"],
            "cls": cat[1] if len(cat) > 1 else "",
            "sector": cat[2] if len(cat) > 2 else "",
            "issued_h": issued_h,
            # الميلادي المقابل لـ 1/1 الاصطلاحي مضلِّل — يُترك فارغًا حين تسجّل الهيئة السنة فقط.
            "issued_g": None if year_only else to_date(law.get("issued_gregorian")),
            "best_h": inst_h if year_only and inst_h else issued_h,
            "date_note": date_note,
            "published_h": fmt_hijri(law.get("published_hijri")),
            "inst_kind": KIND.get(first["kind"], "") if first else "",
            "inst_text": first["text"].strip() if first else "",
            "inst_date": inst_h,
            "others": "؛ ".join(i["text"].strip() for i in insts[1:]),
            "status": STATUS.get(law.get("status_at_source") or "", law.get("status_at_source") or ""),
            "articles": law.get("article_count"),
            "in_hakeem": "نعم" if hid else "",
            "hakeem_name": hakeem_name.get(hid, "") if hid else "",
            "url": law["url"],
        })
    cls_order = {"أنظمة أساسية": 0, "أنظمة عادية": 1, "تنظيمات، وترتيبات تنظيمية": 2, "لوائح وما في حكمها": 3}
    rows.sort(key=lambda r: (cls_order.get(r["cls"], 9), r["issued_g"] or date(1, 1, 1)), reverse=False)

    # ── الورقة 2: أنظمة حكيم بلا مقابل في الهيئة (من الإصدار السابق) ──
    matched_names = {norm(r["hakeem_name"]) for r in rows if r["hakeem_name"]} | {norm(r["name"]) for r in rows}
    unmatched = []
    if PREV.exists():
        prev = load_workbook(PREV, read_only=True)["الأنظمة"]
        header = None
        for vals in prev.iter_rows(values_only=True):
            if header is None:
                header = list(vals)
                continue
            rec = dict(zip(header, vals))
            nm = rec.get("اسم النظام / الأداة") or ""
            if nm and norm(nm) not in matched_names:
                unmatched.append(rec)

    wb = Workbook()
    thin = Side(style="thin", color="D0D7DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E79")
    repealed_fill = PatternFill("solid", fgColor="FDE2E2")

    def sheet(ws, headers, data, widths, repealed_col=None):
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill, c.border = head_font, head_fill, border
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for vals in data:
            ws.append(vals)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.border = border
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(c.value, date):
                    c.number_format = "yyyy-mm-dd"
            if repealed_col is not None and row[repealed_col].value == "ملغى":
                for c in row:
                    c.fill = repealed_fill
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions

    ws1 = wb.active
    ws1.title = "الأنظمة"
    h1 = ["م", "اسم النظام / الأداة (هيئة الخبراء)", "الفئة", "القطاع", "تاريخ الإصدار المعتمد (هـ)",
          "تاريخ الإصدار في الهيئة (هـ)", "ما يوافقه (م)", "ملاحظة التاريخ",
          "تاريخ النشر (هـ)", "نوع أداة الإصدار", "أداة الإصدار", "تاريخ الأداة (هـ)", "أدوات أخرى", "الحالة في الهيئة",
          "عدد المواد", "في قاعدة حكيم", "الاسم في حكيم", "رابط المصدر الرسمي"]
    sheet(ws1, h1, [[i, r["name"], r["cls"], r["sector"], r["best_h"], r["issued_h"], r["issued_g"], r["date_note"],
                     r["published_h"], r["inst_kind"],
                     r["inst_text"], r["inst_date"], r["others"], r["status"], r["articles"], r["in_hakeem"],
                     r["hakeem_name"], r["url"]] for i, r in enumerate(rows, 1)],
          [5, 42, 18, 20, 14, 14, 13, 30, 14, 16, 42, 14, 30, 12, 9, 10, 36, 30], repealed_col=13)

    ws2 = wb.create_sheet("في حكيم بلا مقابل")
    h2 = ["م", "اسم النظام / الأداة (حكيم)", "النوع", "المجال", "عدد المواد", "تاريخ الإصدار (هـ)", "مصدر التاريخ"]
    sheet(ws2, h2, [[i, r.get("اسم النظام / الأداة"), r.get("النوع"), r.get("المجال"), r.get("عدد المواد"),
                     r.get("تاريخ الإصدار (هـ)") or "", r.get("مصدر التاريخ") or ""] for i, r in enumerate(unmatched, 1)],
          [5, 50, 12, 22, 10, 16, 30])

    ws3 = wb.create_sheet("ملخص")
    cls = Counter(r["cls"] for r in rows)
    st = Counter(r["status"] for r in rows)
    kinds = Counter(r["inst_kind"] for r in rows)
    summary = [
        ("إجمالي وثائق الهيئة", len(rows)),
        *[(f"الفئة — {k}", v) for k, v in cls.most_common()],
        *[(f"الحالة — {k}", v) for k, v in st.most_common()],
        *[(f"أداة الإصدار — {k or 'بلا أداة'}", v) for k, v in kinds.most_common()],
        ("بتاريخ إصدار (هـ)", sum(1 for r in rows if r["issued_h"])),
        ("بتاريخ ميلادي", sum(1 for r in rows if r["issued_g"])),
        ("الهيئة تسجّل السنة فقط (1/1)", sum(1 for r in rows if r["date_note"].startswith("الهيئة تسجّل السنة"))),
        ("— منها اكتمل تاريخها من الأداة", sum(1 for r in rows if r["date_note"].endswith("من الأداة"))),
        ("تاريخ الأداة يختلف عن تاريخ الإصدار المسجّل", sum(1 for r in rows if r["date_note"].startswith("تاريخ الأداة"))),
        ("مطابقة لنظام في حكيم (هوية مُتحقَّقة)", sum(1 for r in rows if r["in_hakeem"])),
        ("أنظمة حكيم بلا مقابل في الهيئة", len(unmatched)),
        ("تاريخ رصد الهيئة", "2026-09-28"),
        ("تاريخ التصدير", date.today().isoformat()),
        ("ملاحظة", "التواريخ والأدوات والحالة كما وردت في بوابة هيئة الخبراء حرفيًا؛ لا تاريخ مُضاف من غيرها."),
        ("ملاحظة", "«تاريخ الإصدار المعتمد» = تاريخ الهيئة، إلا حين تسجّل الهيئة السنة فقط (1/1) فيُؤخذ اليوم والشهر من نصّ أداة الإصدار."),
        ("ملاحظة", "المطابقة مع حكيم من مطابقات الهوية المُتحقَّقة (416)؛ غير المطابَق لا يعني غيابه عن حكيم."),
    ]
    sheet(ws3, ["البند", "القيمة"], summary, [48, 70])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"{OUT.relative_to(ROOT)}: {len(rows)} وثيقة، {len(unmatched)} بلا مقابل")


if __name__ == "__main__":
    main()
