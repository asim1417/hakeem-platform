"""ملف قرارات المالك لمراجعة ملاحظات التعديل (من data/review/proposals.jsonl). لا يكتب على أي قاعدة."""
import json
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path(__file__).resolve().parents[1]
rows = [json.loads(l) for l in open(ROOT / "data" / "review" / "proposals.jsonl", encoding="utf-8")]
OUT = ROOT / "reports" / "مراجعة_ملاحظات_التعديل.xlsx"

CATS = {
    1: ("الفئة 1 - تعدد", "أكثر من أداة أو بند في ملاحظة واحدة"),
    2: ("الفئة 2 - بلا أداة", "الأداة غير مقروءة في الملاحظة"),
    3: ("الفئة 3 - بلا نص جديد", "تعديل لا يذكر النص الجديد"),
    4: ("الفئة 4 - أثر غير صريح", "غير مصنف، ملاحظة عامة، جدول، دمج أو ترتيب"),
}
HEAD = ["معرّف الملاحظة", "النظام", "المادة", "الفئة", "الأثر المقترح", "الأداة", "تاريخها", "النص الجديد (حرفي)",
        "الدليل (اقتباس)", "المصدر", "الثقة", "سبب عدم الحسم", "قرار المالك", "ملاحظته",
        "رابط مرفق الأداة", "صفحة الصورة", "صورة الصفحة", "مسودة آلية من الصورة (للمطابقة فقط)", "أساس التصنيف", "تصنيف المحلل الآلي",
        "قراءة الصورة", "الأداة كما في الصورة", "التحقق المستقل", "تصحيحات التحقق", "نص في الصورة لأداة أو مادة أخرى", "ملاحظة القارئ"]
WIDTH = {"معرّف الملاحظة": 30, "النظام": 34, "المادة": 22, "الفئة": 8, "الأثر المقترح": 22, "الأداة": 26, "تاريخها": 12,
         "النص الجديد (حرفي)": 60, "الدليل (اقتباس)": 60, "المصدر": 30, "الثقة": 16, "سبب عدم الحسم": 40, "قرار المالك": 16,
         "ملاحظته": 30, "رابط مرفق الأداة": 30, "صورة الصفحة": 30, "مسودة آلية من الصورة (للمطابقة فقط)": 60,
         "أساس التصنيف": 26, "تصنيف المحلل الآلي": 16, "صفحة الصورة": 10, "قراءة الصورة": 14, "الأداة كما في الصورة": 26,
         "التحقق المستقل": 20, "تصحيحات التحقق": 30, "نص في الصورة لأداة أو مادة أخرى": 50, "ملاحظة القارئ": 40}
F = Font(name="Arial", size=10)
FB = Font(name="Arial", size=10, bold=True)
YELLOW = PatternFill("solid", fgColor="FFF2CC")
GREY = PatternFill("solid", fgColor="E7E6E6")
CONF_FILL = {"مؤكد": PatternFill("solid", fgColor="E2EFDA"), "مرجّح": PatternFill("solid", fgColor="DDEBF7"),
             "يحتاج قرار المالك": PatternFill("solid", fgColor="FCE4D6")}
CLIP = 32000  # حد خلية Excel


def cell(v):
    if v is None:
        return None
    v = str(v)
    return v if len(v) <= CLIP else v[:CLIP] + " …[مقتطع؛ النص كاملًا في proposals.jsonl]"


wb = Workbook()
ws0 = wb.active
ws0.title = "ملخص"
ws0.sheet_view.rightToLeft = True

sheets = {}
for c, (title, _) in CATS.items():
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.append(HEAD)
    for x in ws[1]:
        x.font = FB
        x.fill = GREY
        x.alignment = Alignment(wrap_text=True, vertical="center")
    dv = DataValidation(type="list", formula1='"موافق,مرفوض,موافق مع تصحيح,مؤجل"', allow_blank=True)
    ws.add_data_validation(dv)
    sheets[c] = (ws, dv)

order = {"يحتاج قرار المالك": 0, "مرجّح": 1, "مؤكد": 2}
for r in sorted(rows, key=lambda r: (r["category"], order[r["confidence"]], r["law"], r["unit"] or "")):
    ws, dv = sheets[r["category"]]
    ws.append([cell(x) for x in [
        r["note_id"], r["law"], r["unit"], r["category"], r["effect"], r["instrument"], r["instrument_date"], r["new_text"],
        r["evidence"], r["source"], r["confidence"], r["reason"], None, None,
        r["attachment_url"], r.get("draft_page") or r.get("ocr_page"), r.get("draft_image") or r.get("page_image"),
        r.get("draft_ocr") or r.get("ocr_excerpt"), r["effect_basis"], r["old_effect"],
        {"found": "وُجد", "partial": "جزئي", "not_found": "لم يوجد", "not_applicable": "لا يعدّل النص", "wrong_article": "مادة أخرى"}.get(r.get("vision_status"), r.get("vision_status")),
        r.get("instrument_seen"), r.get("vision_verify") if r.get("vision_status") else None, r.get("vision_fixes"), r.get("vision_other_text"), r.get("vision_notes")]])
    i = ws.max_row
    dv.add(f"M{i}")

for c, (ws, _) in sheets.items():
    for i, h in enumerate(HEAD, start=1):
        col = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[col].width = WIDTH[h]
    for row in ws.iter_rows(min_row=2):
        for x in row:
            x.font = F
            x.alignment = Alignment(wrap_text=True, vertical="top")
        row[12].fill = YELLOW
        row[13].fill = YELLOW
        row[10].fill = CONF_FILL.get(row[10].value, GREY)
        for k in (9, 14, 16):
            if row[k].value and str(row[k].value).startswith("http"):
                row[k].hyperlink = row[k].value
    for x in ws.iter_rows(min_row=2, max_col=HEAD.index("ملاحظته") + 1, min_col=HEAD.index("قرار المالك") + 1):
        pass
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

# ── الملخص (معادلات تحسب من أوراق الفئات) ──
ws0.column_dimensions["A"].width = 44
for col in "BCDEF":
    ws0.column_dimensions[col].width = 16
ws0.append(["مراجعة ملاحظات التعديل غير المحسومة — هيئة الخبراء (جولة 28 سبتمبر 2026)"])
ws0["A1"].font = Font(name="Arial", size=12, bold=True)
ws0.append(["اقتراح فقط: لا يُطبَّق شيء على حكيم قبل قرار المالك. القرارات المعتمدة تُستورد صفوفًا جديدة في uqn_fix.amendment_review."])
ws0.append([])
ws0.append(["الفئة", "مؤكد", "مرجّح", "يحتاج قرار المالك", "المجموع", "وصف الفئة"])
for x in ws0[4]:
    x.font = FB
    x.fill = GREY
first = ws0.max_row + 1
for c, (title, desc) in CATS.items():
    r = ws0.max_row + 1
    q = f"'{title}'!$K:$K"
    ws0.append([title, f'=COUNTIF({q},"مؤكد")', f'=COUNTIF({q},"مرجّح")', f'=COUNTIF({q},"يحتاج قرار المالك")', f"=SUM(B{r}:D{r})", desc])
last = ws0.max_row
ws0.append(["المجموع", f"=SUM(B{first}:B{last})", f"=SUM(C{first}:C{last})", f"=SUM(D{first}:D{last})", f"=SUM(E{first}:E{last})"])
tot = ws0.max_row
for x in ws0[tot]:
    x.font = FB
ws0.append([])
notes_n = len({r["note_id"].split("#")[0] for r in rows})
ws0.append(["عدد الملاحظات المراجَعة", notes_n, "", "", "", "تتفكك الملاحظة المتعددة الآثار إلى صفوف (معرّف#رقم)"])
ws0.append(["صفوف لها مسودة آلية من صورة الأداة", f"=COUNTIF('{CATS[1][0]}'!$R:$R,\"?*\")+COUNTIF('{CATS[2][0]}'!$R:$R,\"?*\")+COUNTIF('{CATS[3][0]}'!$R:$R,\"?*\")+COUNTIF('{CATS[4][0]}'!$R:$R,\"?*\")",
            "", "", "", "المسودة للمطابقة بالصورة فقط؛ الأرقام والتواريخ فيها غير موثوقة"])
ws0.append(["قرارات المالك المدخلة", f"=COUNTA('{CATS[1][0]}'!$M:$M)+COUNTA('{CATS[2][0]}'!$M:$M)+COUNTA('{CATS[3][0]}'!$M:$M)+COUNTA('{CATS[4][0]}'!$M:$M)-4"])
ws0.append([])
ws0.append(["دقة التصنيف الآلي (عينة 20 صفًا مقروءة يدويًا بعد التصحيح)", "19 من 20", "", "", "", "الخطأ الوحيد (ترتيب مع دمج) صُحّحت قاعدته"])
ws0.append(["الفئة 3: صفوف قُرئت صورها بصريًا", 166, "", "", "", "69 وثيقة، 356 صفحة"])
ws0.append(["نصوص تحقق منها قارئ مستقل كلمةً كلمة", 135, "", "", "", "مطابق 119، فروق شكل حروف 10، فروق كلمات 6؛ صُحّح كل ما ثبت بالصورة"])
ws0.append(["مرفق الهيئة أداة غير الأداة المذكورة في الملاحظة", f"=COUNTIF('{CATS[3][0]}'!$L:$L,\"مرفق الهيئة لهذه المادة أداة أخرى*\")", "", "", "", "لا يُنسب نص الأداة الأخرى لأداة الملاحظة"])
ws0.append([])
ws0.append(["طريقة الاستعمال"])
ws0[f"A{ws0.max_row}"].font = FB
for line in [
    "العمودان الأصفران «قرار المالك» و«ملاحظته» لك وحدك. القيم: موافق / مرفوض / موافق مع تصحيح / مؤجل.",
    "«مؤكد»: الأثر صريح والأداة مقروءة والنص الجديد منقول حرفيًا من ملاحظة الهيئة (أو الأثر لا يحتاج نصًا كالحذف والإلغاء).",
    "«مرجّح»: الأثر مستنتج بقاعدة (دمج، ترتيب، إجراء إداري، آثار متعددة) أو الأداة مستنتجة من ملاحظات النظام نفسه.",
    "«يحتاج قرار المالك»: نقص في الدليل، وسببه في عمود «سبب عدم الحسم».",
    "الفئة 3: النص الجديد ليس في الملاحظة، فنُقل من صورة الأداة قراءةً بصرية وتحقق منه قارئ ثانٍ؛ افتح «صورة الصفحة» وطابقه قبل «موافق».",
    "إذا كان مرفق الهيئة أداة أخرى، بقي الصف «يحتاج قرار المالك» ونص تلك الأداة في عمود «نص في الصورة لأداة أو مادة أخرى».",
    "صورة الهيئة للمادة تعرض النص الأصلي عند الإصدار؛ النص المعدَّل لا يُؤخذ منها.",
    "مثال: موافق مع تصحيح | «النص الصحيح بعد المطابقة: …».",
]:
    ws0.append([line])
for row in ws0.iter_rows():
    for x in row:
        if x.font is None or not x.font.bold:
            x.font = Font(name="Arial", size=10, bold=bool(x.font and x.font.bold))
        x.alignment = Alignment(wrap_text=True, vertical="top")
ws0["A1"].font = Font(name="Arial", size=12, bold=True)

OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print(OUT, len(rows))
