import json, re, collections, datetime, copy, shutil
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from hijridate import Hijri, Gregorian

REPO = '/home/claude/hakeem-platform/ingest/uqn-fix-2026-09-25/'
B = REPO + 'data/'
SRC_XLSX = REPO + 'مراجعة_التعديلات_والإحلالات.xlsx'
OUT1 = '/home/claude/work/out/قائمة-قرارات-حقن-أم-القرى.xlsx'
OUT2 = '/home/claude/work/out/مراجعة_التعديلات_والإحلالات — مع السريان المعاد.xlsx'
TODAY = datetime.date(2026, 10, 18)

def h2g(s):
    y, m, d = map(int, s.split('-')); g = Hijri(y, m, d).to_gregorian()
    return datetime.date(g.year, g.month, g.day)
def g2h(g):
    h = Gregorian(g.year, g.month, g.day).to_hijri(); return f"{h.year}-{h.month:02d}-{h.day:02d}"
def addd(h, n): return g2h(h2g(h) + datetime.timedelta(days=n))
def norm(h):
    m = re.search(r'(1[34]\d\d)\D(\d{1,2})\D(\d{1,2})', str(h or ''))
    return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}" if m else None

ops = [json.loads(l) for l in open(B + 'amendment_operations.jsonl')]
sup = [json.loads(l) for l in open(B + 'supersessions_verified.jsonl')]
man = [json.loads(l) for l in open(B + 'amendment_manual_queue.jsonl')]
scan = json.load(open('/home/claude/work/scan.json'))
live = {x['id']: x for x in (json.loads(l) for l in open('/home/claude/work/live.jsonl') if l.strip())}

# sheet membership
wb0 = openpyxl.load_workbook(SRC_XLSX)
sheet_of = {}
for sh in ('مراجعة سريعة', 'مراجعة دقيقة'):
    for r in wb0[sh].iter_rows(min_row=2, values_only=True):
        if r[3]: sheet_of[r[3]] = sh

KIND = {'royal_decree': 'مرسوم ملكي', 'cabinet_decision': 'قرار مجلس الوزراء', 'royal_order': 'أمر ملكي'}
OPAR = {'replace_law': 'إحلال نظام', 'repeal_law': 'إلغاء نظام', 'insert_paragraph': 'إضافة فقرة', 'repeal_article': 'إلغاء مادة',
        'replace_article_text': 'تعديل نص مادة', 'replace_paragraph_text': 'تعديل نص فقرة', 'append_text': 'إضافة نص',
        'delete_paragraph': 'حذف/إلغاء فقرة', 'substitute_phrase': 'استبدال عبارة', 'replace_definition': 'تعديل تعريف',
        'replace_definition_entry': 'تعديل بند تعريف', 'rename_law': 'تعديل اسم نظام'}
ART71 = 'سكوت الأداة + المادة ٧١ من النظام الأساسي للحكم («…وتكون نافذة المفعول من تاريخ نشرها، ما لم ينص على تاريخ آخر») — متحقق بأداة hakeem_verify_citation'
OLD_RULE = 'من تاريخ النشر (افتراض — لم يُنص على مهلة في الأداة)'

# supersession map by old title → record (live-verified)
S_LIVE = {0: 'S00n', 1: 'S01n', 2: 'S02n', 3: 'S03n', 4: 'S04n', 5: 'S05n', 6: 'S06n', 7: 'S07n', 8: 'S08n', 9: 'S08n',
          10: 'S10n', 11: 'S10n', 12: 'S12n', 13: 'S13n', 14: 'S14n', 15: 'S15n', 16: 'S16n'}
S_DAYS = {0: 180, 1: 180, 2: 180, 3: 120, 4: 180, 5: 180, 6: 180, 7: 90, 8: 180, 9: 180, 10: 90, 11: 90}
sup_rows = []
for j, s in enumerate(sup):
    lv = live.get(S_LIVE[j], {})
    pub_live = norm(lv.get('pub_hijri'))
    q = '؛ '.join(lv.get('effective_verbatim') or [])
    if j in S_DAYS:
        comp = addd(pub_live, S_DAYS[j]); rule = f'النشر ({pub_live}) + {S_DAYS[j]} يومًا'
        if j == 7: rule += ' — والإحلال «عند نفاذه» (البند ثانياً من المرسوم)'
        if j in (10, 11): rule += ' — والإلغاء «اعتباراً من تاريخ نفاذ» النظام (قرار م.و ٦٨)'
    elif j in (12, 14, 15):
        comp = pub_live; rule = f'من تاريخ النشر ({pub_live}) بنص التنظيم'
    elif j == 13:
        comp = '1448-02-28'; rule = 'من تاريخ الموافقة: قرار مجلس الوزراء (219) وتاريخ 1448/02/28هـ — نص التنظيم «يعمل بالتنظيم من تاريخ الموافقة عليه»؛ لا من النشر (' + pub_live + ')'
    else:
        comp = None; rule = 'غير مثبت: لا نص إلغاء أو إحلال في قرار الموافقة (628/1443)؛ التنظيم الجديد نفسه يسري من نشره ' + str(pub_live)
    match = (comp == s['effective_from_hijri'])
    sup_rows.append(dict(j=j, s=s, pub_live=pub_live, comp=comp, rule=rule, quote=q, live_url=s['new_source_url'], match=match,
                         greg=str(h2g(comp)) if comp else None))

def sup_for(title):
    for r in sup_rows:
        if r['s']['old_title'] == title and r['comp']: return r

# ---------- per-instrument effective determination ----------
inst = collections.OrderedDict()
for o in ops:
    inst.setdefault((o['instrument_kind'], o['instrument_no'], o['source_url']), []).append(o)
idx_of = {tuple(v['key']): int(i) for i, v in scan.items()}

def decide(i, o, pub):
    """return (date_hijri or None, sanad, status, quote) status in silent/explicit/linked/conditional/failed"""
    lid = f"I{i:03d}"; lv = live[lid]
    t = o['op_type']; title = o['target_law_title']
    if i == 8:
        return addd_year(pub), 'البند ثانياً من المرسوم م/87: نفاذ التعديل بمضي سنة من النشر (' + pub + ')؛ السنة هجرية بالتقويم الرسمي', 'explicit', lv['effective_verbatim'][0]
    if i == 1:
        d = addd(pub, 180)
        return d, ('الإلغاء في البند ثالثاً مرتبط بالنظام الجديد الذي «يُعمل بالنظام بعد (مائة وثمانين) يومًا من تاريخ نشره» (المادة الأخيرة، صفحة النظام https://www.uqn.gov.sa/4001465) — النشر ' + pub +
                   ' + 180؛ مع استمرار العمل بأحكام النظام الملغى حتى إقرار مجلس شؤون التعليم العام (خلال سنة من النفاذ)'), 'linked', (lv.get('linked_effective_verbatim') or [''])[0]
    if i == 9 and t == 'delete_paragraph':
        r = sup_for('نظام تأديب الموظفين')
        return r['comp'], 'البند سادساً: الإلغاء «اعتباراً من تاريخ نفاذ نظام هيئة الرقابة ومكافحة الفساد» = نشره ' + r['pub_live'] + ' + 90 يومًا (المادة الأخيرة من النظام)؛ عدا المادة 47', 'linked', lv['effective_verbatim'][0]
    if i == 85:
        r = sup_for('تنظيم الهيئة الوطنية لمكافحة الفساد')
        return r['comp'], 'البند تاسعاً: الإلغاء «اعتباراً من تاريخ نفاذ النظام» = نشر نظام هيئة الرقابة ' + r['pub_live'] + ' + 90 يومًا', 'linked', lv['effective_verbatim'][1]
    if i in (57, 77):
        r = sup_for('نظام تملك غير السعوديين للعقار واستثماره')
        extra = '' if i == 57 else ' — تنبيه: البند نصه «اتخاذ الإجراءات النظامية لإلغاء العمل…» أي توجيه بإجراء لا إلغاء ذاتي النفاذ'
        return r['comp'], 'مرتبط بنفاذ نظام تملك غير السعوديين للعقار (م/14) = نشره ' + r['pub_live'] + ' + 180 يومًا (المادة الأخيرة)' + extra, 'linked', lv['effective_verbatim'][-1]
    if i == 32:
        return None, 'البند رابعاً: يسري التعديلان «في تاريخ متزامن مع تاريخ نفاذ» نظام الإثبات — نص نظام الإثبات الختامي لم يُقرأ (الصفحة لا تربط به ولم يُعثر عليه بالبحث)', 'failed', lv['effective_verbatim'][0]
    if i == 138:
        return '1443-05-28', 'البند رابعاً: تاريخ محدد بالنص (لا مهلة)', 'explicit', lv['effective_verbatim'][0]
    if i == 154:
        return g2h(datetime.date(2021, 4, 1)), 'البند أولاً: الإلغاء «اعتباراً من تاريخ 1/4/2021م» (تاريخ ميلادي محدد بالنص، يقابله ' + g2h(datetime.date(2021, 4, 1)) + 'هـ)', 'explicit', lv['effective_verbatim'][0]
    if i == 155:
        return None, 'مشروط: «يعمل بما ورد في البند (ثانياً) … اعتباراً من تاريخ استكمال الإجراءات المنصوص عليها في البند (أولاً)» — لا تاريخ منشور', 'conditional', lv['effective_verbatim'][0]
    if i == 153 and t == 'replace_article_text':
        return None, 'مشروط: «ويبدأ العمل بهذا البند بعد انتهاء مدة عضوية الأعضاء المعينين بقرار مجلس الوزراء رقم (213)» — يُتحقق هل العبارة من صلب القرار أم من النص المعدَّل', 'conditional', lv['effective_verbatim'][0]
    if i == 102:
        return pub, 'البند ثالثاً: «ويعمل بها من تاريخ نشرها في الجريدة الرسمية» (نص صريح)', 'explicit', lv['effective_verbatim'][0]
    if t == 'replace_law' and i == 158:
        return pub, ART71 + ' — الإحلال هنا إحلال جهة (وزارة البيئة) محل هيئة في الاختصاصات، لا نظام جديد', 'silent', ''
    if t == 'replace_law':
        r = sup_for(title)
        if r:
            return r['comp'], 'من المواد الختامية للتنظيم الجديد: ' + r['rule'], 'linked', r['quote']
        return None, 'إحلال: تاريخ السريان في المادة الختامية للتنظيم/النظام الجديد، ونصه «بالصيغة المرافقة» لم يُنشر في صفحة الأداة ولم يُقرأ', 'failed', ''
    if i == 9:
        return pub, ART71 + ' — تنبيه: المرسوم يعتمد نظام هيئة الرقابة الذي يسري بعد 90 يومًا؛ الأرجح أن يُربط الإحلال بنفاذه، ويحتاج قرار المالك', 'silent', ''
    return pub, ART71, 'silent', ''

def addd_year(h):
    y, m, d = map(int, h.split('-')); return f"{y+1}-{m:02d}-{d:02d}"

JUD = ['نظام القضاء', 'المرافعات الشرعية', 'نظام التنفيذ', 'نظام الإثبات', 'نظام التحكيم', 'نظام المحاماة', 'الإجراءات الجزائية',
       'ديوان المظالم', 'الأحوال الشخصية', 'المعاملات المدنية', 'المحاكم التجارية']
def is_jud(title):
    if any(k in title for k in JUD) and 'اللائحة التنفيذية' not in title: return True
    return title.strip() in ('نظام العمل',)
PEN = re.compile(r'(عقوب|غرام|السجن|بالسجن|يعاقب|يُعاقب|جزاء|الجزاءات)')
PEN_TITLE = re.compile(r'(مكافحة|الأسلحة|المخدرات|الإرهاب|الرشوة|التحرش)')

groups = []
op_result = {}
for gi, (k, v) in enumerate(inst.items(), 1):
    i = idx_of[k]; lv = live[f"I{i:03d}"]; pub = norm(lv['pub_hijri'])
    assert pub == v[0]['published_hijri']
    res = []
    for o in v:
        d, sanad, st, q = decide(i, o, pub)
        changed = (d != o['effective_from_hijri'])
        op_result[o['op_id']] = dict(date=d, sanad=sanad, st=st, q=q, changed=changed, url=k[2], pub=pub)
        res.append((o, d, sanad, st, q, changed))
    sheets = {sheet_of.get(o['op_id'], '?') for o in v}
    titles = sorted({o['target_law_title'] for o in v})
    jud = any(is_jud(t) for t in titles)
    pen = any(PEN.search((o.get('evidence_quote') or '')) for o in v) or any(PEN_TITLE.search(t) for t in titles)
    sts = {r[3] for r in res}
    reasons = []
    if jud: reasons.append('يمس نظامًا قضائيًّا')
    if pen: reasons.append('عقوبات/غرامات أو نظام جزائي')
    if sheets != {'مراجعة سريعة'}: reasons.append('فيه عمليات «مراجعة دقيقة»')
    if sts & {'failed', 'conditional'}: reasons.append('تاريخ السريان تعذّر/مشروط')
    level = 'جماعي' if not reasons else 'فردي'
    dates = sorted({r[1] or '—' for r in res})
    sanads = list(dict.fromkeys(r[2] for r in res))
    if 'failed' in sts: prop = 'أجّل الحقن: اقرأ نص التنظيم/النظام المرتبط أولًا (انظر ورقة «تعذّر»)'
    elif 'conditional' in sts: prop = 'اعتمد النص بلا تاريخ نفاذ آلي: يُحقن «صادر لم يسرِ» حتى يثبت تحقق الشرط'
    elif level == 'جماعي': prop = 'اعتمد جماعيًّا بتاريخ السريان المعاد'
    else: prop = 'راجع النصوص فرديًّا ثم اعتمد بتاريخ السريان المعاد'
    future = [x for x in dates if x != '—' and h2g(x) > TODAY]
    if future: prop += ' — تنبيه: السريان لم يحن بعد (' + '، '.join(future) + ')'
    groups.append(dict(no=gi, i=i, kind=KIND[k[0]], num=k[1], date=v[0]['instrument_date_hijri'], pub=pub, url=k[2], titles=titles,
                       n=len(v), ids=[o['op_id'] for o in v], types=collections.Counter(OPAR[o['op_type']] for o in v),
                       eff=dates, sanad=sanads, level=level, reasons=reasons, prop=prop, sheets=sheets,
                       changed=sum(r[5] for r in res), sts=sts))

json.dump(op_result, open('/home/claude/work/op_result.json', 'w'), ensure_ascii=False, indent=0)

# ---------- styles ----------
F = Font(name='Arial', size=10); FB = Font(name='Arial', size=10, bold=True, color='FFFFFF')
HFILL = PatternFill('solid', fgColor='1F4E78'); YEL = PatternFill('solid', fgColor='FFF2CC'); RED = PatternFill('solid', fgColor='F8CBAD')
GRN = PatternFill('solid', fgColor='E2EFDA'); ORG = PatternFill('solid', fgColor='FCE4D6')
WR = Alignment(wrap_text=True, vertical='top', horizontal='right', readingOrder=2)
thin = Side(style='thin', color='BFBFBF'); BOR = Border(left=thin, right=thin, top=thin, bottom=thin)

def sheet(wb, name, headers, widths, rows, fills=None):
    ws = wb.create_sheet(name); ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for c in ws[1]: c.font = FB; c.fill = HFILL; c.alignment = WR; c.border = BOR
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = F; c.alignment = WR; c.border = BOR
    for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    return ws

wb = openpyxl.Workbook(); wb.remove(wb.active)

# Summary sheet
nG = len(groups); nJ = sum(g['level'] == 'جماعي' for g in groups)
nchg = sum(1 for r in op_result.values() if r['changed'] and r['date'])
fails = [g for g in groups if g['sts'] & {'failed'}]; conds = [g for g in groups if g['sts'] & {'conditional'}]
ws = wb.create_sheet('ملخص'); ws.sheet_view.rightToLeft = True
lines = [
    ('قائمة قرارات حقن تعديلات وإحلالات أم القرى', ''),
    ('أُعدّت', '2026-10-18 (تشغيل مجدول — قراءة وإعداد فقط؛ لم يُكتب على أي قاعدة ولم يُعبأ عمود «القرار»)'),
    ('المصدر', 'asim1417/hakeem-platform — ingest/uqn-fix-2026-09-25 (395 عملية من 159 أداة) + قراءة حية لكل صفحة بأداة WebFetch (175 صفحة، 0 فشل في الصفحات الأساسية)'),
    ('عدد المجموعات', nG), ('جماعي', nJ), ('فردي', nG - nJ),
    ('عمليات تغيّر تاريخ سريانها عن الافتراض', nchg),
    ('عمليات سندها صار «سكوت الأداة + م٧١» بدل «افتراض»', sum(r['st'] == 'silent' for r in op_result.values())),
    ('مجموعات فيها تاريخ تعذّر تحديده', len(fails)), ('مجموعات فيها سريان مشروط', len(conds)),
    ('قاعدة الحساب', 'المدة تُحسب من تاريخ النشر الظاهر في صفحة أم القرى (لا من تاريخ المرسوم أو القرار)؛ الأيام تُضاف تقويميًّا ثم تُحوّل بتقويم أم القرى (مكتبة hijridate)؛ «سنة» = سنة هجرية. اصطلاح العدّ: النشر + N يومًا (يوم النشر هو اليوم صفر) — مطابق لما في الحزمة.'),
    ('المادة ٧١', 'نصها كما أرجعته hakeem_verify_citation: «تنشر الأنظمة في الجريدة الرسمية، وتكون نافذة المفعول من تاريخ نشرها، ما لم ينص على تاريخ آخر.» (الحكم: الإحالة صحيحة — الحالة قيد التحقق)'),
    ('مستوى المراجعة', '«جماعي» = كل عمليات المجموعة من «مراجعة سريعة» وخارج النطاق القضائي وبلا عقوبات وتاريخها محدد. وأُضيف شرط: ما تعذّر أو اشترط تاريخه يُنقل إلى «فردي».'),
    ('عمود «قرار المالك»', 'فارغ (أصفر) في ورقة «القرارات» — للتعبئة: اعتمد / ارفض / وريني النصوص'),
]
for a, b in lines: ws.append([a, b])
for row in ws.iter_rows():
    for c in row: c.font = F; c.alignment = WR
ws['A1'].font = Font(name='Arial', size=13, bold=True)
ws.column_dimensions['A'].width = 38; ws.column_dimensions['B'].width = 120

def fmt_types(c): return '، '.join(f'{k} ({v})' for k, v in c.most_common())
rows = []
for g in groups:
    rows.append([g['no'], f"{g['kind']} ({g['num'] or 'بلا رقم'}) وتاريخ {g['date']}هـ — نُشر {g['pub']}هـ", ' | '.join(g['titles']), g['n'], '، '.join(g['ids']),
                 fmt_types(g['types']), '، '.join(g['eff']) + ('' if not g['changed'] else f"  (تغيّر في {g['changed']})"),
                 ' || '.join(g['sanad']), g['level'], '؛ '.join(g['reasons']) or '—', g['prop'], '', g['url']])
ws = sheet(wb, 'القرارات', ['رقم المجموعة', 'الأداة ورقمها وتاريخها', 'الأنظمة المستهدفة', 'عدد العمليات', 'معرّفاتها', 'أنواعها',
                            'تاريخ السريان المعاد (هـ)', 'سنده', 'مستوى المراجعة', 'سبب المستوى', 'المقترح', 'قرار المالك', 'رابط المصدر'],
           [8, 34, 34, 8, 30, 26, 18, 60, 10, 26, 36, 14, 34], rows)
for r in range(2, ws.max_row + 1):
    ws.cell(r, 12).fill = YEL
    lv = ws.cell(r, 9); lv.fill = GRN if lv.value == 'جماعي' else ORG
    if 'تغيّر' in str(ws.cell(r, 7).value): ws.cell(r, 7).fill = RED
ws.cell(1, 12).fill = PatternFill('solid', fgColor='BF8F00')
# order: put individual after? keep instrument order; add filter.

# Supersessions sheet
rows = []
order = sorted(sup_rows, key=lambda r: (r['j'] != 0, r['j']))
for r in order:
    s = r['s']
    now = '—' if not r['comp'] else ('نعم' if h2g(r['comp']) <= TODAY else f"لا — يسري بعد {(h2g(r['comp'])-TODAY).days} يومًا")
    rows.append([r['j'] + 1, s['old_title'], s.get('old_hakeem_code'), s.get('old_instrument'), s['new_title'], {'replace': 'إحلال', 'repeal': 'إلغاء', 'repeal_with_saving': 'إلغاء مع بقاء أحكام'}[s['relation']],
                 s['effective_from_hijri'], s['effective_rule'], r['pub_live'], r['comp'], r['greg'], r['rule'], r['quote'],
                 'مطابق' if r['match'] else ('غير مثبت' if not r['comp'] else 'مختلف'), now, s['new_source_url'], s.get('notes')])
ws = sheet(wb, 'الإحلالات', ['#', 'النظام القديم', 'رمز حكيم', 'سند القديم', 'النظام الجديد', 'العلاقة', 'يسري من (الحزمة)', 'قاعدة الحزمة',
                             'تاريخ النشر (قراءة حية)', 'يسري من — متحقق (هـ)', 'يقابله (م)', 'سند السريان', 'النص الحرفي الدال (WebFetch)', 'المطابقة', 'ساري اليوم 2026-10-18؟', 'رابط النظام الجديد', 'ملاحظة الحزمة'],
           [5, 30, 10, 20, 30, 12, 12, 18, 13, 13, 12, 50, 50, 10, 16, 34, 34], rows)
ws.cell(2, 15).fill = RED; ws.cell(2, 10).fill = RED

# Evidence sheet: one row per instrument
rows = []
for g in groups:
    lv = live[f"I{g['i']:03d}"]
    q = '؛ '.join(lv.get('effective_verbatim') or []) or 'لا نص سريان (سكوت)'
    rows.append([g['no'], g['kind'], g['num'], g['date'], g['url'], lv.get('title'), g['pub'], lv.get('pub_greg'),
                 'رأس صفحة أم القرى (WebFetch حي 2026-10-18)، مطابق للقطة الحزمة', q, '، '.join(g['eff']), ' || '.join(g['sanad']),
                 '، '.join(sorted(g['sts'])), (lv.get('notes') or '')[:600]])
for r in order:
    if r['comp'] or r['j'] == 16:
        rows.append([f"إحلال {r['j']+1}", 'نظام/تنظيم جديد', '', '', r['live_url'], r['s']['new_title'], r['pub_live'], live[S_LIVE[r['j']]].get('pub_greg'),
                     'رأس صفحة النظام في أم القرى (WebFetch حي)', r['quote'], r['comp'] or '—', r['rule'], 'linked', ''])
ws = sheet(wb, 'السريان', ['المجموعة', 'نوع الأداة', 'رقمها', 'تاريخها (هـ)', 'الرابط', 'عنوان الصفحة', 'تاريخ النشر (هـ)', 'تاريخ النشر (م)',
                           'من أين قُرئ تاريخ النشر', 'النص الحرفي الدال على السريان', 'التاريخ المحسوب (هـ)', 'السند', 'الحالة', 'ملاحظات القراءة'],
           [8, 14, 9, 11, 34, 34, 11, 11, 26, 60, 14, 60, 10, 50], rows)

# Failed sheet
rows = []
for g in groups:
    for oid in g['ids']:
        r = op_result[oid]
        if r['st'] in ('failed', 'conditional'):
            rows.append([g['no'], oid, f"{g['kind']} ({g['num']})", ' | '.join(g['titles']), 'تعذّر' if r['st'] == 'failed' else 'مشروط', r['sanad'], r['q'], g['url']])
rows.append(['—', '—', 'روابط تابعة لم تُفتح', '', 'تعذّر', 'روابط نصوص مرافقة أعادت الصفحة الرئيسية أو CLIENT_ERROR: ?p=20038، ?p=17247، ?p=13757، uqn.gov.sa/2021/04/23/g2/، صفحة نظام الاحتيال المالي (م/79)، صفحة إلغاء مكتب براءات الخليج (م/71)، ?p=16013', '', ''])
rows.append(['—', '—', 'إحلال 17 (مدينة الملك عبدالله للطاقة الذرية)', '', 'غير مثبت', 'لا نص إلغاء أو إحلال في قرار 628/1443 ولا في التنظيم (قراءة حية) — لا يُضاف سجل تحقق', '', 'https://www.uqn.gov.sa/details?p=19806'])
ws = sheet(wb, 'تعذّر', ['المجموعة', 'المعرّف', 'الأداة', 'النظام', 'الحالة', 'السبب', 'النص المقروء', 'الرابط'], [8, 16, 20, 34, 9, 70, 50, 34], rows)

# Manual sheet summary
rows = []
for m in man:
    tags = ' '.join(m['unparsed_items'])
    if 'بالصيغة المرافقة' in tags and 'ملحق' in tags: why = 'ملحق «بالصيغة المرافقة» غير منشور في الصفحة'
    elif 'البند يعدّل قرارًا' in tags: why = 'يعدّل قرارًا أو أمرًا لا نص نظام'
    else: why = 'صيغة لم يغطها المستخرج / بنود متعددة أو وصفية'
    rows.append([m['effect_id'], KIND.get(m['instrument_kind'], m['instrument_kind']), m['instrument_no'], m['instrument_date_hijri'], m['target_title_as_cited'], why, len(m['unparsed_items']), m['source_url']])
rows.sort(key=lambda r: (r[5], r[4]))
ws = sheet(wb, 'يدوي — ما ينقص', ['رقم الأثر', 'الأداة', 'رقمها', 'تاريخها', 'النظام كما ورد', 'ما ينقص', 'عدد البنود', 'الرابط'], [11, 14, 9, 11, 40, 40, 9, 40], rows)
man_counts = collections.Counter(r[5] for r in rows)

wb.move_sheet('القرارات', offset=-wb.index(wb['القرارات']) + 1)
wb.save(OUT1)

# ---------- copy of original with two added columns ----------
shutil.copy(SRC_XLSX, OUT2)
wb2 = openpyxl.load_workbook(OUT2)
for sh in ('مراجعة سريعة', 'مراجعة دقيقة'):
    ws = wb2[sh]
    hdr = [c.value for c in ws[1]]; assert hdr[11] == 'يسري من (هـ)' and hdr[12] == 'قاعدة السريان', hdr
    before = [[c.value for c in row] for row in ws.iter_rows()]
    widths = {k: v.width for k, v in ws.column_dimensions.items()}
    ws.insert_cols(14, 2)
    for k in list(ws.column_dimensions.keys()):
        pass
    for col_letter, w in widths.items():
        j = openpyxl.utils.column_index_from_string(col_letter)
        nj = j + 2 if j >= 14 else j
        ws.column_dimensions[get_column_letter(nj)].width = w
    ws.column_dimensions['N'].width = 13; ws.column_dimensions['O'].width = 50
    h = ws.cell(1, 13)
    for j, txt in ((14, 'يسري من — معاد'), (15, 'سند السريان')):
        c = ws.cell(1, j, txt); c.font = copy.copy(h.font); c.fill = copy.copy(h.fill); c.alignment = copy.copy(h.alignment); c.border = copy.copy(h.border)
    for r in range(2, ws.max_row + 1):
        oid = ws.cell(r, 4).value
        if not oid: continue
        res = op_result[oid]
        ws.cell(r, 14, res['date'] or ('مشروط' if res['st'] == 'conditional' else 'تعذّر'))
        ws.cell(r, 15, res['sanad'])
        for j in (14, 15):
            c = ws.cell(r, j); c.font = Font(name='Arial', size=10); c.alignment = Alignment(wrap_text=True, vertical='top', readingOrder=2)
            if res['changed']: c.fill = RED
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    after = [[c.value for c in row] for row in ws.iter_rows()]
    for b, a in zip(before, after):
        assert b[:13] == a[:13] and b[13:] == a[15:], 'existing cell changed'
    assert all(ws.cell(r, 1).value is None for r in range(2, ws.max_row + 1))
wb2.save(OUT2)

# stats for message
stats = dict(nG=nG, nJ=nJ, nF=nG - nJ, nchg=nchg, fails=[(g['no'], g['kind'], g['num'], g['titles']) for g in fails],
             conds=[(g['no'], g['kind'], g['num'], g['titles']) for g in conds], man=man_counts,
             changed=[(g['no'], g['kind'], g['num'], g['titles'], g['eff'], g['pub']) for g in groups if g['changed']],
             sup_mismatch=[r['j'] for r in sup_rows if r['comp'] and not r['match']])
json.dump(stats, open('/home/claude/work/stats.json', 'w'), ensure_ascii=False, indent=1)
print(json.dumps(stats, ensure_ascii=False, indent=1)[:6000])
