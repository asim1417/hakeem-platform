import json,re,pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
R=json.load(open('attach.json'))
BAD={'تنظيم مكتبة الملك فهد الوطنية','تنظيم مركز الإيرادات غير النفطية','الترتيبات التنظيمية للمعهد الوطني لأبحاث الصحة'}
for u,v in R.items():
    if v['law'] in BAD: v['cm']=None
AR=str.maketrans('٠١٢٣٤٥٦٧٨٩','0123456789')
def strip_head(t):
    t=re.sub(r'^تم نسخ الرابط\s*','',t); t=re.sub(r'^#[^\n]*\n','',t.strip()); t=re.sub(r'^\d{4}-\d{1,2}-\d{1,2}\nالموافق\n\d{2}-\d{2}-\d{4}\n','',t.strip())
    return t.strip()
def dec_meta(title):
    m=re.search(r'\((م\s*/\s*[\d٠-٩]+)\)\s*وتاريخ\s*([\d٠-٩/‏ ]+)هـ',title.translate(AR))
    return (m.group(1).replace(' ',''),re.sub(r'[\s‏]','',m.group(2))) if m else ('','')
def cm_kind(t,title,u):
    h=t[:400]
    if '/news/' in u or 'رأس خادم الحرمين' in h or 'واس' in h[:200]: return 'خبر جلسة مجلس الوزراء (ليس نص القرار)'
    if re.search(r'أمر ملكي|أمر كريم',h): return 'أمر ملكي'
    if re.search(r'برقية تعميمية|تعميم رقم',h): return 'تعميم الديوان الملكي'
    if re.search(r'قرار\s+رقم',h): return 'قرار مجلس الوزراء'
    return 'أداة أخرى'
def cm_meta(t):
    m=re.search(r'(?:برقية تعميمية|تعميم)\s+رقم\s*\(\s*([\d٠-٩]+)\s*\)\s*وتاريخ\s*([\d٠-٩/‏ ]+)هـ',t.translate(AR))
    if m: return (m.group(1),re.sub(r'[\s‏]','',m.group(2)))
    m=re.search(r'قرار\s+رقم\s*\(\s*([\d٠-٩]+)\s*\)\s*وتاريخ\s*([\d٠-٩/‏ ]+)هـ',t.translate(AR))
    return (m.group(1),re.sub(r'[\s‏]','',m.group(2))) if m else ('','')
def key_clauses(t):
    out=[]
    for s in t.split('\n'):
        if re.search(r'(يحل|إحلال|إلغاء|يلغي|تلغى|حذف|تعديل)',s) and re.match(r'^\s*(أولاً|ثانياً|ثالثاً|رابعاً|خامساً|سادساً|سابعاً|ثامناً|تاسعاً|عاشراً|أولًا|ثانيًا|ثالثًا|رابعًا|خامسًا|سادسًا|سابعًا|ثامنًا|تاسعًا|حادي|ثاني)',s): out.append(s.strip())
    return out
L15=json.load(open('articles15.json')); L66=json.load(open('articles66.json'))
def build(L,tag):
    rows=[];laws=[];js=[]
    for l in L:
        v=R[l['url']]; pre=[]
        dno=ddt=cno=cdt=''; clauses=[]
        if v['decree']:
            u,title,t=v['decree']; dno,ddt=dec_meta(title); body=strip_head(t)
            pre.append({'الرقم':0,'التسمية':'المرسوم الملكي','العنوان':title,'النص':body,'المصدر':u}); clauses+=key_clauses(body)
        if v['cm']:
            u,title,t=v['cm']; body=strip_head(t); cno,cdt=cm_meta(body[:500]); 
            kind=cm_kind(body,title,u); pre.append({'الرقم':0,'التسمية':kind,'العنوان':title,'النص':body,'المصدر':u}); clauses+=key_clauses(body)
        eff=''
        if l['articles']:
            last=l['articles'][-1]['text']; m=re.search(r'(يُ?عمل\s+ب[^.\n]*)',last); eff=m.group(1) if m else ''
        for p in pre: rows.append({'النظام':l['law'],'الموضع':'أدوات الإصدار',**p,'الباب':'','الفصل':''})
        for a in l['articles']:
            rows.append({'النظام':l['law'],'الموضع':l.get('mode','مواد'),'الرقم':a['n'],'التسمية':a['label'],'العنوان':'','النص':a['text'],'المصدر':l['url'],'الباب':a.get('bab',''),'الفصل':a.get('fasl','')})
        laws.append({'النظام':l['law'],'رقم المرسوم':dno,'تاريخ المرسوم':ddt,'أداة الموافقة':(cm_kind(strip_head(v['cm'][2]),v['cm'][1],v['cm'][0]) if v['cm'] else ''),'رقمها':cno,'تاريخها':cdt,'النفاذ':eff,'أحكام الإحلال والإلغاء والتعديل في أدوات الإصدار':'\n'.join(dict.fromkeys(clauses)),'عدد المواد':len(l['articles']),'أدوات الإصدار':('مرسوم + قرار' if v['decree'] and v['cm'] else 'مرسوم فقط' if v['decree'] else 'قرار فقط' if v['cm'] else 'غير موجودة'),'ملاحظة':(v['cm'][1] if v['cm'] and 'خبر' in v['cm'][1] else '')+(' '+l.get('note','') if l.get('note') else ''),'رابط النص':l['url']})
        js.append({'law':l['law'],'issuance':{'royal_decree':{'no':dno,'date_h':ddt,'text':pre[0]['النص'] if v['decree'] else '','url':v['decree'][0] if v['decree'] else ''},'cabinet_decision':{'no':cno,'date_h':cdt,'text':(pre[-1]['النص'] if v['cm'] else ''),'url':v['cm'][0] if v['cm'] else ''}},'effective':eff,'source_url':l['url'],'articles':[{'number':a['n'],'label':a['label'],'chapter':a.get('bab',''),'section':a.get('fasl',''),'text':a['text']} for a in l['articles']]})
    return pd.DataFrame(laws),pd.DataFrame(rows),js
for L,name in [(L15+L66,'الأنظمة_الواحد_والثمانون_مع_أدوات_الإصدار')]:
    S,A,js=build(L,name)
    print(S['أدوات الإصدار'].value_counts().to_dict(), len(A))
    out=name+'.xlsx'
    with pd.ExcelWriter(out) as w:
        S.to_excel(w,sheet_name='الأنظمة',index=False); A.to_excel(w,sheet_name='أدوات الإصدار والمواد',index=False)
    wb=load_workbook(out)
    W={'النظام':32,'النص':90,'أحكام الإحلال والإلغاء والتعديل في أدوات الإصدار':70,'ملاحظة':40,'المصدر':38,'رابط النص':38,'العنوان':34,'النفاذ':40,'التسمية':20,'الباب':22,'الفصل':22}
    for ws in wb:
        ws.sheet_view.rightToLeft=True
        for c in ws[1]: c.font=Font(name='Arial',bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F3A5F'); c.alignment=Alignment(horizontal='center',wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for c in row: c.font=Font(name='Arial'); c.alignment=Alignment(horizontal='right',vertical='top',wrap_text=True)
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=W.get(col[0].value,12)
        ws.freeze_panes='B2'
    wb.save(out)
    json.dump(js,open(name+'.json','w'),ensure_ascii=False,indent=1)
    print(S[S['أدوات الإصدار']=='غير موجودة']['النظام'].tolist())
    print('')
